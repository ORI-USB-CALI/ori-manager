"""Exportación del convenio a Excel (.xlsx) con openpyxl (BE-03 de HU-24).

Parte del mismo DTO que el PDF, así ambos formatos salen con los mismos datos.
"""

from __future__ import annotations

from collections.abc import Sequence
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from backend.schemas.exportacion import (
    INFORMACION_NO_DISPONIBLE,
    MARCA_BORRADOR,
    ConvenioPdfDTO,
    ModoEmision,
)
from backend.services.pdf_convenio import ServicioPdfConvenio

HOJA_DATOS_GENERALES = "Datos Generales"
HOJA_PARTICIPANTES_COMPROMISOS = "Participantes y Compromisos"

_FUENTE_CABECERA = Font(bold=True)
_FONDO_CABECERA = PatternFill(fill_type="solid", start_color="D9D9D9", end_color="D9D9D9")
_FUENTE_TITULO_SECCION = Font(bold=True, size=12)
_FUENTE_MARCA_BORRADOR = Font(bold=True, size=12, color="B91C1C")
_AJUSTE_TEXTO = Alignment(wrap_text=True, vertical="top")
_ANCHO_MINIMO = 10
_ANCHO_MAXIMO = 80

_ETIQUETAS_SECCION = {"heading": "Sección", "paragraph": "Cláusula", "list_item": "Ítem"}

# Mismo criterio de seguridad que el reporte de HU-34 (OWASP CSV/Excel
# injection): un texto que empieza por uno de estos caracteres se guarda con
# un apóstrofo delante para que la hoja de cálculo no lo evalúe como fórmula.
_PREFIJOS_FORMULA = ("=", "+", "-", "@")


def _texto_seguro(valor: object) -> object:
    if isinstance(valor, str) and valor.startswith(_PREFIJOS_FORMULA):
        return f"'{valor}"
    return valor


class ConvenioExcelExportService:
    """Genera el libro .xlsx del convenio en memoria."""

    def generar_excel(self, convenio_id: int, db: Session) -> bytes:
        dto = ServicioPdfConvenio().mapear_convenio_a_dto(convenio_id, db)
        return self.construir_libro(dto)

    def construir_libro(self, dto: ConvenioPdfDTO) -> bytes:
        libro = Workbook()
        datos_generales = libro.active
        datos_generales.title = HOJA_DATOS_GENERALES
        self._escribir_datos_generales(datos_generales, dto)
        self._escribir_participantes_y_compromisos(
            libro.create_sheet(HOJA_PARTICIPANTES_COMPROMISOS), dto
        )
        for hoja in libro.worksheets:
            _ajustar_columnas(hoja)

        salida = BytesIO()
        libro.save(salida)
        return salida.getvalue()

    def _escribir_datos_generales(self, hoja: Worksheet, dto: ConvenioPdfDTO) -> None:
        inicio = _escribir_marca(hoja, dto)
        _escribir_cabecera(hoja, inicio, ["Campo", "Valor"])
        etiqueta_identificacion, valor_identificacion = dto.identificacion()
        filas: list[tuple[str, object]] = [
            (etiqueta_identificacion, valor_identificacion),
            ("Consecutivo de la solicitud", dto.solicitud_consecutivo),
            ("Tipo de convenio", dto.tipo_convenio),
            ("Estado", dto.estado),
            ("Etapa", dto.etapa),
            ("Alcance", dto.alcance),
            ("Objeto", dto.objeto),
            ("Justificación", dto.justificacion),
            ("Fecha de inicio de vigencia", dto.fecha_inicio),
            ("Fecha de fin de vigencia", dto.fecha_vencimiento),
            ("Duración (meses)", dto.duracion_meses),
            ("Porcentaje de avance", dto.porcentaje_avance),
            ("Versión del documento", dto.version_actual),
            ("Unidad solicitante", dto.solicitante_unidad),
        ]
        for fila, (campo, valor) in enumerate(filas, start=inicio + 1):
            hoja.cell(row=fila, column=1, value=campo).font = _FUENTE_CABECERA
            # Un dato ausente se declara como tal, nunca se supone.
            _escribir_valor(
                hoja, fila, 2, INFORMACION_NO_DISPONIBLE if valor is None else valor
            )
        hoja.freeze_panes = f"A{inicio + 1}"

    def _escribir_participantes_y_compromisos(
        self, hoja: Worksheet, dto: ConvenioPdfDTO
    ) -> None:
        fila = _escribir_tabla(
            hoja,
            _escribir_marca(hoja, dto),
            "Participantes",
            [
                "Parte",
                "Nombre",
                "Tipo",
                "Identificación",
                "Unidad",
                "Ciudad",
                "País",
                "Contacto",
                "Correo",
            ],
            [
                [
                    "Parte solicitante",
                    dto.solicitante_nombre,
                    dto.solicitante_tipo,
                    None,
                    dto.solicitante_unidad,
                    None,
                    None,
                    dto.solicitante_nombre,
                    dto.solicitante_correo,
                ],
                [
                    "Entidad aliada",
                    dto.aliado_nombre,
                    dto.aliado_tipo,
                    dto.aliado_identificacion,
                    None,
                    dto.aliado_ciudad,
                    dto.aliado_pais,
                    dto.aliado_contacto_nombre,
                    dto.aliado_contacto_correo,
                ],
            ],
        )
        fila = _escribir_tabla(
            hoja,
            fila + 1,
            "Firmantes",
            ["Orden", "Rol", "Parte", "Nombre", "Cargo", "Modalidad", "Estado"],
            [
                [f.orden, f.rol, f.parte, f.nombre, f.cargo, f.modalidad, f.estado]
                for f in dto.firmantes
            ]
            or [["Sin firmantes registrados"]],
        )
        compromisos: list[list[object]] = [
            ["Implicación financiera", dto.implicacion_financiera]
        ]
        compromisos.extend(
            [_ETIQUETAS_SECCION.get(seccion.tipo, seccion.tipo), seccion.texto]
            for seccion in dto.secciones_contenido
        )
        _escribir_tabla(hoja, fila + 1, "Compromisos", ["Tipo", "Detalle"], compromisos)


def _escribir_marca(hoja: Worksheet, dto: ConvenioPdfDTO) -> int:
    """En un borrador, la marca va en la primera fila; devuelve la fila libre."""
    if dto.modo_emision is not ModoEmision.BORRADOR:
        return 1
    hoja.cell(row=1, column=1, value=MARCA_BORRADOR).font = _FUENTE_MARCA_BORRADOR
    return 3


def _escribir_cabecera(hoja: Worksheet, fila: int, titulos: Sequence[str]) -> None:
    for columna, titulo in enumerate(titulos, start=1):
        celda = hoja.cell(row=fila, column=columna, value=titulo)
        celda.font = _FUENTE_CABECERA
        celda.fill = _FONDO_CABECERA


def _escribir_valor(hoja: Worksheet, fila: int, columna: int, valor: object) -> None:
    celda = hoja.cell(row=fila, column=columna, value=_texto_seguro(valor))
    if isinstance(valor, str) and len(valor) > _ANCHO_MAXIMO:
        celda.alignment = _AJUSTE_TEXTO


def _escribir_tabla(
    hoja: Worksheet,
    fila: int,
    titulo: str,
    cabecera: Sequence[str],
    filas: Sequence[Sequence[object]],
) -> int:
    """Escribe título, cabecera y filas; devuelve la siguiente fila libre."""
    hoja.cell(row=fila, column=1, value=titulo).font = _FUENTE_TITULO_SECCION
    _escribir_cabecera(hoja, fila + 1, cabecera)
    siguiente = fila + 2
    for valores in filas:
        for columna, valor in enumerate(valores, start=1):
            _escribir_valor(hoja, siguiente, columna, valor)
        siguiente += 1
    return siguiente + 1


def _ajustar_columnas(hoja: Worksheet) -> None:
    """Ancho según el contenido más largo de cada columna, con un tope."""
    anchos: dict[int, int] = {}
    for fila in hoja.iter_rows():
        for celda in fila:
            if celda.value is None:
                continue
            largo = max(len(linea) for linea in str(celda.value).splitlines() or [""])
            anchos[celda.column] = max(anchos.get(celda.column, 0), largo)
    for columna, largo in anchos.items():
        hoja.column_dimensions[get_column_letter(columna)].width = min(
            max(largo + 2, _ANCHO_MINIMO), _ANCHO_MAXIMO
        )

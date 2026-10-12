"""HU-34: reporte consolidado de convenios activos (JSON, PDF y Excel)."""

from collections.abc import Callable
from datetime import date
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from backend.core.roles import CodigoRol, TipoUsuario
from backend.models.actividad_convenio import ActividadConvenio
from backend.models.aliado import Aliado
from backend.models.convenio import Convenio
from backend.models.enums import (
    EstadoActividadConvenio,
    EstadoConvenio,
    ParteEjecutoraActividad,
    TipoIdentificacion,
)
from backend.models.unidad_organizacional import UnidadOrganizacional
from backend.models.usuario import Usuario

URL_JSON = "/api/convenios/reportes/activos"
URL_PDF = "/api/convenios/reportes/activos/pdf"
URL_EXCEL = "/api/convenios/reportes/activos/excel"


@pytest.fixture
def aliado_reporte(db: Session) -> Aliado:
    aliado = Aliado(
        nombre="Universidad Aliada HU-34",
        tipo="UNIVERSIDAD",
        tipo_identificacion=TipoIdentificacion.NIT.value,
        identificacion=uuid4().hex,
    )
    db.add(aliado)
    db.commit()
    db.refresh(aliado)
    return aliado


@pytest.fixture
def unidad_reporte(db: Session) -> UnidadOrganizacional:
    unidad = UnidadOrganizacional(
        codigo=f"U-{uuid4().hex[:8]}", nombre="Facultad HU-34", tipo="FACULTAD"
    )
    db.add(unidad)
    db.commit()
    db.refresh(unidad)
    return unidad


@pytest.fixture
def crear_convenio_reporte(
    db: Session, gestor: Usuario, crear_convenio, aliado_reporte, unidad_reporte
) -> Callable[..., Convenio]:
    def _crear(
        *,
        estado: EstadoConvenio,
        codigo: str | None = None,
        fecha_inicio: date | None = date(2026, 1, 1),
        fecha_vencimiento: date | None = date(2027, 1, 1),
        con_aliado_formal: bool = True,
    ) -> Convenio:
        convenio = crear_convenio(
            gestor,
            codigo=codigo,
            fecha_inicio=fecha_inicio,
            fecha_vencimiento=fecha_vencimiento,
        )
        convenio.estado = estado.value
        if con_aliado_formal:
            convenio.aliado_id = aliado_reporte.id
        convenio.unidad_organizacional_id = unidad_reporte.id
        db.commit()
        db.refresh(convenio)
        return convenio

    return _crear


@pytest.fixture
def crear_actividad(db: Session, gestor: Usuario) -> Callable[..., ActividadConvenio]:
    """Crea una `ActividadConvenio` con los valores mínimos requeridos por
    la entidad canónica aprobada (descripcion y registrado_por_id son
    obligatorios)."""

    def _crear(
        convenio: Convenio,
        *,
        titulo: str = "Actividad de prueba",
        descripcion: str = "Descripción de la actividad de prueba",
        estado: EstadoActividadConvenio = EstadoActividadConvenio.REALIZADA,
        parte_ejecutora: ParteEjecutoraActividad | None = ParteEjecutoraActividad.USB,
        fecha_planificada: date | None = None,
        fecha_ejecucion: date | None = date(2026, 3, 2),
        observaciones: str | None = None,
        responsable_nombre: str | None = None,
        responsable_id: int | None = None,
        registrado_por: Usuario | None = None,
    ) -> ActividadConvenio:
        actividad = ActividadConvenio(
            convenio_id=convenio.id,
            titulo=titulo,
            descripcion=descripcion,
            estado=estado.value,
            parte_ejecutora=parte_ejecutora.value if parte_ejecutora else None,
            fecha_planificada=fecha_planificada,
            fecha_ejecucion=fecha_ejecucion,
            observaciones=observaciones,
            responsable_nombre=responsable_nombre,
            responsable_id=responsable_id,
            registrado_por_id=(registrado_por or gestor).id,
        )
        db.add(actividad)
        db.commit()
        db.refresh(actividad)
        return actividad

    return _crear


# ---- CA-01: solo convenios activos -----------------------------------------


def test_incluye_exclusivamente_vigente_y_por_vencer(
    client, gestor, crear_convenio_reporte
) -> None:
    for estado in (EstadoConvenio.VIGENTE, EstadoConvenio.POR_VENCER):
        crear_convenio_reporte(estado=estado, codigo=f"CONV-{estado.value}")

    respuesta = client.get(URL_JSON)

    assert respuesta.status_code == 200
    estados = {fila["estado"] for fila in respuesta.json()}
    assert estados == {"VIGENTE", "POR_VENCER"}


def test_excluye_estados_no_activos_y_heredados(
    client, gestor, crear_convenio_reporte
) -> None:
    for estado in (
        EstadoConvenio.EN_TRAMITE,
        EstadoConvenio.FINALIZADO,
        EstadoConvenio.CANCELADO,
        EstadoConvenio.VENCIDO,
        EstadoConvenio.RENOVADO,
    ):
        crear_convenio_reporte(estado=estado, codigo=f"CONV-{estado.value}")

    respuesta = client.get(URL_JSON)

    assert respuesta.json() == []


# ---- CA-02: informacion principal ------------------------------------------


def test_muestra_informacion_principal_del_convenio(
    client, gestor, crear_convenio_reporte, aliado_reporte, unidad_reporte
) -> None:
    convenio = crear_convenio_reporte(
        estado=EstadoConvenio.VIGENTE, codigo="CONV-INFO"
    )

    respuesta = client.get(URL_JSON)

    assert respuesta.status_code == 200
    (fila,) = respuesta.json()
    assert fila["codigo"] == "CONV-INFO"
    assert fila["aliado"] == aliado_reporte.nombre
    assert fila["unidad_organizacional"] == unidad_reporte.nombre
    assert fila["responsable"] == gestor.nombre_completo
    assert fila["fecha_inicio"] == convenio.fecha_inicio.isoformat()
    assert fila["fecha_vencimiento"] == convenio.fecha_vencimiento.isoformat()


def test_convenio_sin_codigo_no_rompe_el_reporte(
    client, gestor, crear_convenio_reporte
) -> None:
    convenio = crear_convenio_reporte(estado=EstadoConvenio.VIGENTE, codigo=None)

    respuesta = client.get(URL_JSON)

    (fila,) = respuesta.json()
    assert fila["codigo"] is None
    assert fila["id"] == convenio.id


def test_contraparte_propuesta_sin_aliado_formal(
    client, gestor, crear_convenio_reporte
) -> None:
    crear_convenio_reporte(
        estado=EstadoConvenio.VIGENTE, codigo="CONV-PROPUESTA", con_aliado_formal=False
    )

    respuesta = client.get(URL_JSON)

    (fila,) = respuesta.json()
    # El fixture `crear_convenio` (conftest.py) registra la solicitud con
    # nombre_aliado_propuesto="Universidad Contraparte".
    assert fila["aliado"] == "Universidad Contraparte"


# ---- CA-03: convenios activos sin actividades ------------------------------


def test_convenio_activo_sin_actividades_aparece_con_lista_vacia(
    client, gestor, crear_convenio_reporte
) -> None:
    crear_convenio_reporte(estado=EstadoConvenio.VIGENTE, codigo="CONV-SIN-ACT")

    respuesta = client.get(URL_JSON)

    (fila,) = respuesta.json()
    assert fila["actividades"] == []


def test_convenio_activo_con_varias_actividades_las_incluye(
    client, db, gestor, crear_convenio_reporte, crear_actividad
) -> None:
    convenio = crear_convenio_reporte(estado=EstadoConvenio.VIGENTE, codigo="CONV-ACT")
    crear_actividad(
        convenio,
        titulo="Taller de movilidad académica",
        parte_ejecutora=ParteEjecutoraActividad.USB,
        estado=EstadoActividadConvenio.REALIZADA,
        fecha_planificada=date(2026, 3, 1),
        fecha_ejecucion=date(2026, 3, 2),
    )
    crear_actividad(
        convenio,
        titulo="Encuentro de seguimiento",
        parte_ejecutora=ParteEjecutoraActividad.AMBAS,
        estado=EstadoActividadConvenio.PENDIENTE,
        fecha_ejecucion=None,
    )

    respuesta = client.get(URL_JSON)

    (fila,) = respuesta.json()
    assert len(fila["actividades"]) == 2
    titulos = {actividad["titulo"] for actividad in fila["actividades"]}
    assert titulos == {"Taller de movilidad académica", "Encuentro de seguimiento"}


# ---- CA-04: informacion actualizada -----------------------------------------


def test_refleja_cambios_posteriores_del_convenio(
    client, db, gestor, crear_convenio_reporte
) -> None:
    convenio = crear_convenio_reporte(estado=EstadoConvenio.VIGENTE, codigo="CONV-UPD")
    assert client.get(URL_JSON).json()[0]["objeto"] != "Objeto actualizado"

    convenio.objeto = "Objeto actualizado"
    db.commit()

    respuesta = client.get(URL_JSON)
    assert respuesta.json()[0]["objeto"] == "Objeto actualizado"


def test_refleja_actividades_registradas_despues_de_la_primera_consulta(
    client, db, gestor, crear_convenio_reporte, crear_actividad
) -> None:
    convenio = crear_convenio_reporte(estado=EstadoConvenio.VIGENTE, codigo="CONV-REGEN")
    assert client.get(URL_JSON).json()[0]["actividades"] == []

    crear_actividad(convenio, titulo="Actividad registrada después")
    # El fixture `db` usa expire_on_commit=False: la colección ya cargada en
    # el identity map no se invalida sola tras el commit.
    db.expire_all()

    respuesta = client.get(URL_JSON)
    assert len(respuesta.json()[0]["actividades"]) == 1


# ---- Modelo compartido con HU-35 -------------------------------------------


def test_realizada_sin_fecha_ejecucion_viola_la_restriccion_de_bd(
    db: Session, gestor, crear_convenio_reporte
) -> None:
    """La BD debe impedir una actividad REALIZADA sin fecha_ejecucion,
    sin importar qué historia la registre."""
    convenio = crear_convenio_reporte(estado=EstadoConvenio.VIGENTE, codigo="CONV-CHECK")
    db.add(
        ActividadConvenio(
            convenio_id=convenio.id,
            titulo="Actividad inválida",
            descripcion="Sin fecha de ejecución",
            estado=EstadoActividadConvenio.REALIZADA.value,
            fecha_ejecucion=None,
            registrado_por_id=gestor.id,
        )
    )
    with pytest.raises(IntegrityError):
        db.commit()
    db.rollback()


def test_reporte_lee_actividades_con_la_estructura_que_usara_hu_35(
    client, gestor, crear_convenio_reporte, crear_actividad
) -> None:
    """Simula cómo HU-35 registrará actividades en la entidad compartida:
    siempre REALIZADA, sin parte_ejecutora ni fecha_planificada, con
    responsable en texto libre y quien registra identificado aparte. El
    reporte de HU-34 debe leerlas sin transformaciones ni tabla adicional.
    """
    convenio = crear_convenio_reporte(estado=EstadoConvenio.VIGENTE, codigo="CONV-HU35")
    crear_actividad(
        convenio,
        titulo="Visita de seguimiento",
        descripcion="Visita registrada como lo hará HU-35",
        estado=EstadoActividadConvenio.REALIZADA,
        parte_ejecutora=None,
        fecha_planificada=None,
        fecha_ejecucion=date(2026, 4, 10),
        observaciones="Sin novedades",
        responsable_nombre="Persona externa sin cuenta en el sistema",
        registrado_por=gestor,
    )

    respuesta = client.get(URL_JSON)

    (fila,) = respuesta.json()
    assert len(fila["actividades"]) == 1
    assert fila["actividades"][0]["titulo"] == "Visita de seguimiento"
    assert fila["actividades"][0]["estado"] == "REALIZADA"
    assert fila["actividades"][0]["fecha_ejecucion"] == "2026-04-10"


# ---- Permisos ---------------------------------------------------------------


def test_requiere_autenticacion(client) -> None:
    assert client.get(URL_JSON).status_code == 401


def test_gestor_puede_consultar(client, gestor) -> None:
    assert client.get(URL_JSON).status_code == 200


def test_administrador_puede_consultar(client, crear_usuario, entrar_como) -> None:
    entrar_como(crear_usuario(CodigoRol.ADMINISTRADOR_ORI))
    assert client.get(URL_JSON).status_code == 200


@pytest.mark.parametrize(
    ("rol", "tipo"),
    [
        (CodigoRol.REVISOR_ORI, TipoUsuario.INTERNO),
        (CodigoRol.SOLICITANTE_INTERNO, TipoUsuario.INTERNO),
        (CodigoRol.SOLICITANTE_EXTERNO, TipoUsuario.EXTERNO),
    ],
)
def test_otros_roles_no_pueden_consultar(
    client, crear_usuario, entrar_como, rol, tipo
) -> None:
    entrar_como(crear_usuario(rol, tipo))
    assert client.get(URL_JSON).status_code == 403


# ---- Descargas PDF ------------------------------------------------------------


def test_descarga_pdf_con_convenios(client, gestor, crear_convenio_reporte) -> None:
    crear_convenio_reporte(estado=EstadoConvenio.VIGENTE, codigo="CONV-PDF")

    respuesta = client.get(URL_PDF)

    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"] == "application/pdf"
    assert "attachment" in respuesta.headers["content-disposition"]
    assert respuesta.content.startswith(b"%PDF")
    # reportlab no comprime el stream por defecto: el texto real queda
    # legible dentro de los bytes, así que se valida contenido real y no
    # solo encabezados/magic bytes.
    assert b"CONV-PDF" in respuesta.content


def test_descarga_pdf_incluye_identificacion_sin_codigo(
    client, gestor, crear_convenio_reporte
) -> None:
    convenio = crear_convenio_reporte(estado=EstadoConvenio.VIGENTE, codigo=None)

    respuesta = client.get(URL_PDF)

    assert f"Convenio #{convenio.id}".encode() in respuesta.content


def test_descarga_pdf_sin_convenios_no_falla(client, gestor) -> None:
    respuesta = client.get(URL_PDF)

    assert respuesta.status_code == 200
    assert respuesta.content.startswith(b"%PDF")
    assert b"No hay convenios activos" in respuesta.content


def test_descarga_pdf_con_muchos_convenios_actividades_y_textos_extensos(
    client, db, gestor, crear_convenio_reporte, crear_actividad
) -> None:
    texto_largo = "Objeto extenso. " * 60
    for i in range(12):
        convenio = crear_convenio_reporte(
            estado=EstadoConvenio.VIGENTE, codigo=f"CONV-MASIVO-{i}"
        )
        convenio.objeto = texto_largo
        db.commit()
        crear_actividad(
            convenio,
            titulo="Actividad con descripción & caracteres <especiales>",
            parte_ejecutora=ParteEjecutoraActividad.CONTRAPARTE,
        )

    respuesta = client.get(URL_PDF)

    assert respuesta.status_code == 200
    assert respuesta.content.startswith(b"%PDF")


# ---- Descargas Excel ----------------------------------------------------------


def test_descarga_excel_con_convenios_y_actividades(
    client, db, gestor, crear_convenio_reporte, crear_actividad
) -> None:
    convenio = crear_convenio_reporte(estado=EstadoConvenio.VIGENTE, codigo="CONV-XLS")
    crear_actividad(
        convenio,
        titulo="Encuentro académico",
        parte_ejecutora=ParteEjecutoraActividad.AMBAS,
        estado=EstadoActividadConvenio.PENDIENTE,
        fecha_ejecucion=None,
    )

    respuesta = client.get(URL_EXCEL)

    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in respuesta.headers["content-disposition"]

    libro = load_workbook(BytesIO(respuesta.content))
    hoja_convenios = libro["Convenios activos"]
    fila_convenio = [celda.value for celda in hoja_convenios[2]]
    assert fila_convenio[0] == "CONV-XLS"
    assert fila_convenio[-1] == 1  # N.º de actividades

    hoja_actividades = libro["Actividades"]
    fila_actividad = [celda.value for celda in hoja_actividades[2]]
    assert fila_actividad[0] == "CONV-XLS"
    assert fila_actividad[1] == "Encuentro académico"


def test_descarga_excel_identifica_convenios_sin_codigo(
    client, gestor, crear_convenio_reporte
) -> None:
    convenio = crear_convenio_reporte(estado=EstadoConvenio.VIGENTE, codigo=None)

    respuesta = client.get(URL_EXCEL)

    libro = load_workbook(BytesIO(respuesta.content))
    hoja_convenios = libro["Convenios activos"]
    fila_convenio = [celda.value for celda in hoja_convenios[2]]
    assert fila_convenio[0] == f"Convenio #{convenio.id}"


def test_descarga_excel_sanitiza_textos_que_parecen_formulas(
    client, db, gestor, crear_convenio_reporte, crear_actividad
) -> None:
    """OWASP CSV/Excel injection: un objeto/título que empiece con
    =, +, - o @ no debe exportarse como fórmula ejecutable."""
    convenio = crear_convenio_reporte(estado=EstadoConvenio.VIGENTE, codigo="CONV-INJ")
    convenio.objeto = "=cmd|'/c calc'!A0"
    db.commit()
    crear_actividad(
        convenio,
        titulo="+SUM(1+1)",
        descripcion="Actividad con título que parece fórmula",
    )

    respuesta = client.get(URL_EXCEL)

    libro = load_workbook(BytesIO(respuesta.content))
    hoja_convenios = libro["Convenios activos"]
    celda_objeto = hoja_convenios[2][1]
    assert celda_objeto.value == "'=cmd|'/c calc'!A0"
    assert celda_objeto.data_type != "f"

    hoja_actividades = libro["Actividades"]
    celda_titulo = hoja_actividades[2][1]
    assert celda_titulo.value == "'+SUM(1+1)"
    assert celda_titulo.data_type != "f"


def test_descarga_excel_requiere_permiso(client, crear_usuario, entrar_como) -> None:
    entrar_como(crear_usuario(CodigoRol.REVISOR_ORI))
    assert client.get(URL_EXCEL).status_code == 403


def test_descarga_pdf_requiere_permiso(client, crear_usuario, entrar_como) -> None:
    entrar_como(crear_usuario(CodigoRol.REVISOR_ORI))
    assert client.get(URL_PDF).status_code == 403

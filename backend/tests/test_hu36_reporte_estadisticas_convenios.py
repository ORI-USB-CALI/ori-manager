"""HU-36: reporte estadístico de convenios (JSON, PDF y Excel)."""

import re
from collections.abc import Callable
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import event
from sqlalchemy.orm import Session

from backend.core.roles import CodigoRol, TipoUsuario
from backend.models.aliado import Aliado
from backend.models.convenio import Convenio
from backend.models.enums import EstadoConvenio, TipoAliado
from backend.models.pais import Pais
from backend.models.usuario import Usuario
from backend.services.estadisticas_convenios import (
    Conteo,
    ReporteEstadisticoConvenios,
    ServicioEstadisticasConvenios,
    generar_excel_estadisticas_convenios,
)

URL_JSON = "/api/convenios/reportes/estadisticas"
URL_PDF = "/api/convenios/reportes/estadisticas/pdf"
URL_EXCEL = "/api/convenios/reportes/estadisticas/excel"


@pytest.fixture
def crear_pais(db: Session) -> Callable[[str], Pais]:
    def _crear(nombre: str) -> Pais:
        pais = Pais(codigo_iso=uuid4().hex[:2].upper(), nombre=nombre)
        db.add(pais)
        db.commit()
        db.refresh(pais)
        return pais

    return _crear


@pytest.fixture
def crear_aliado(db: Session) -> Callable[..., Aliado]:
    def _crear(
        *,
        tipo: TipoAliado = TipoAliado.UNIVERSIDAD,
        activo: bool = True,
        pais: Pais | None = None,
    ) -> Aliado:
        aliado = Aliado(
            nombre=f"Aliado {uuid4().hex[:8]}",
            tipo=tipo.value,
            sector_economico="Educación" if tipo == TipoAliado.EMPRESA else None,
            tipo_identificacion="NIT",
            identificacion=uuid4().hex,
            pais_id=pais.id if pais else None,
            activo=activo,
        )
        db.add(aliado)
        db.commit()
        db.refresh(aliado)
        return aliado

    return _crear


@pytest.fixture
def crear_convenio_estadistica(
    db: Session, gestor: Usuario, crear_convenio
) -> Callable[..., Convenio]:
    def _crear(
        *,
        estado: EstadoConvenio,
        aliado: Aliado | None = None,
        nombre_propuesto: str | None = None,
        tipo_propuesto: TipoAliado | None = None,
        pais_propuesto: str | None = None,
    ) -> Convenio:
        convenio = crear_convenio(gestor)
        convenio.estado = estado.value
        convenio.aliado_id = aliado.id if aliado else None
        convenio.solicitud.nombre_aliado_propuesto = nombre_propuesto
        convenio.solicitud.tipo_aliado_propuesto = tipo_propuesto.value if tipo_propuesto else None
        convenio.solicitud.pais_aliado_propuesto = pais_propuesto
        db.commit()
        db.refresh(convenio)
        return convenio

    return _crear


# ---- CA-01/CA-02: agrupaciones con datos reales ----------------------------


def test_agrupa_por_estado_contractual(
    client, gestor, crear_convenio_estadistica
) -> None:
    crear_convenio_estadistica(estado=EstadoConvenio.VIGENTE)
    crear_convenio_estadistica(estado=EstadoConvenio.VIGENTE)
    crear_convenio_estadistica(estado=EstadoConvenio.FINALIZADO)

    respuesta = client.get(URL_JSON)

    assert respuesta.status_code == 200
    datos = respuesta.json()
    assert datos["datos_suficientes"] is True
    assert datos["total_convenios"] == 3
    por_estado = {c["etiqueta"]: c["cantidad"] for c in datos["por_estado_contractual"]}
    assert por_estado["VIGENTE"] == 2
    assert por_estado["FINALIZADO"] == 1


def test_agrupa_por_vigencia_activa(
    client, gestor, crear_convenio_estadistica
) -> None:
    crear_convenio_estadistica(estado=EstadoConvenio.VIGENTE)
    crear_convenio_estadistica(estado=EstadoConvenio.POR_VENCER)
    crear_convenio_estadistica(estado=EstadoConvenio.VENCIDO)
    crear_convenio_estadistica(estado=EstadoConvenio.CANCELADO)

    respuesta = client.get(URL_JSON)

    por_vigencia = {c["etiqueta"]: c["cantidad"] for c in respuesta.json()["por_vigencia_activa"]}
    assert por_vigencia["Vigencia activa"] == 2
    assert por_vigencia["Vigencia no activa"] == 2


def test_agrupa_por_estado_y_tipo_de_aliado(
    client, gestor, crear_convenio_estadistica, crear_aliado
) -> None:
    activo_universidad = crear_aliado(tipo=TipoAliado.UNIVERSIDAD, activo=True)
    inactivo_empresa = crear_aliado(tipo=TipoAliado.EMPRESA, activo=False)
    crear_convenio_estadistica(estado=EstadoConvenio.VIGENTE, aliado=activo_universidad)
    crear_convenio_estadistica(estado=EstadoConvenio.VIGENTE, aliado=inactivo_empresa)

    respuesta = client.get(URL_JSON)

    datos = respuesta.json()
    por_estado_aliado = {c["etiqueta"]: c["cantidad"] for c in datos["por_estado_aliado"]}
    assert por_estado_aliado["Aliado activo"] == 1
    assert por_estado_aliado["Aliado inactivo"] == 1
    por_tipo = {c["etiqueta"]: c["cantidad"] for c in datos["por_tipo_aliado"]}
    assert por_tipo["Universidad"] == 1
    assert por_tipo["Empresa"] == 1


def test_agrupa_por_pais(
    client, gestor, crear_convenio_estadistica, crear_aliado, crear_pais
) -> None:
    colombia = crear_pais("Colombia")
    mexico = crear_pais("México")
    crear_convenio_estadistica(
        estado=EstadoConvenio.VIGENTE, aliado=crear_aliado(pais=colombia)
    )
    crear_convenio_estadistica(
        estado=EstadoConvenio.VIGENTE, aliado=crear_aliado(pais=colombia)
    )
    crear_convenio_estadistica(
        estado=EstadoConvenio.VIGENTE, aliado=crear_aliado(pais=mexico)
    )

    respuesta = client.get(URL_JSON)

    por_pais = {c["etiqueta"]: c["cantidad"] for c in respuesta.json()["por_pais"]}
    assert por_pais["Colombia"] == 2
    assert por_pais["México"] == 1


def test_convenio_sin_aliado_se_agrupa_como_sin_aliado_y_sin_pais(
    client, gestor, crear_convenio_estadistica
) -> None:
    crear_convenio_estadistica(estado=EstadoConvenio.VIGENTE, aliado=None)

    respuesta = client.get(URL_JSON)

    datos = respuesta.json()
    assert datos["por_estado_aliado"] == [{"etiqueta": "Sin aliado registrado", "cantidad": 1}]
    assert datos["por_tipo_aliado"] == [{"etiqueta": "Sin tipo registrado", "cantidad": 1}]
    assert datos["por_pais"] == [{"etiqueta": "Sin país registrado", "cantidad": 1}]
    assert client.get(URL_JSON, params={"pais": "   "}).json()["total_convenios"] == 1


def test_no_duplica_convenios_al_consultar_aliado_y_pais(
    client, gestor, crear_convenio_estadistica, crear_aliado, crear_pais
) -> None:
    """El join con aliado/país no debe multiplicar las filas de convenio."""
    colombia = crear_pais("Colombia")
    crear_convenio_estadistica(
        estado=EstadoConvenio.VIGENTE, aliado=crear_aliado(pais=colombia)
    )

    respuesta = client.get(URL_JSON)

    assert respuesta.json()["total_convenios"] == 1


# ---- Filtros ------------------------------------------------------------------


def test_contraparte_propuesta_en_elaboracion_aporta_tipo_y_pais(
    client, gestor, crear_convenio_estadistica
) -> None:
    convenio = crear_convenio_estadistica(
        estado=EstadoConvenio.EN_TRAMITE,
        nombre_propuesto="Empresa propuesta",
        tipo_propuesto=TipoAliado.EMPRESA,
        pais_propuesto="  España  ",
    )

    datos = client.get(URL_JSON).json()

    assert datos["total_convenios"] == 1
    assert datos["por_estado_contractual"] == [{"etiqueta": "EN_TRAMITE", "cantidad": 1}]
    assert datos["por_vigencia_activa"] == [{"etiqueta": "Vigencia no activa", "cantidad": 1}]
    assert datos["por_estado_aliado"] == [
        {"etiqueta": "Contraparte propuesta (sin formalizar)", "cantidad": 1}
    ]
    assert datos["por_tipo_aliado"] == [{"etiqueta": "Empresa", "cantidad": 1}]
    assert datos["por_pais"] == [{"etiqueta": "España", "cantidad": 1}]
    assert convenio.aliado_id is None
    assert convenio.solicitud.nombre_aliado_propuesto == "Empresa propuesta"
    assert convenio.solicitud.pais_aliado_propuesto == "  España  "


@pytest.mark.parametrize("pais", [None, "   "])
def test_propuesta_sin_tipo_ni_pais_no_inventa_datos(
    client, gestor, crear_convenio_estadistica, pais
) -> None:
    crear_convenio_estadistica(
        estado=EstadoConvenio.EN_TRAMITE,
        nombre_propuesto="Contraparte sin completar",
        pais_propuesto=pais,
    )

    datos = client.get(URL_JSON).json()

    assert datos["por_tipo_aliado"] == [{"etiqueta": "Sin tipo registrado", "cantidad": 1}]
    assert datos["por_pais"] == [{"etiqueta": "Sin país registrado", "cantidad": 1}]


def test_distingue_estados_formales_propuestos_y_sin_informacion(
    client, gestor, crear_convenio_estadistica, crear_aliado
) -> None:
    for activo in (True, False):
        crear_convenio_estadistica(
            estado=EstadoConvenio.VIGENTE, aliado=crear_aliado(activo=activo)
        )
    crear_convenio_estadistica(
        estado=EstadoConvenio.EN_TRAMITE, nombre_propuesto="Contraparte propuesta"
    )
    crear_convenio_estadistica(estado=EstadoConvenio.EN_TRAMITE)

    datos = client.get(URL_JSON).json()

    assert {c["etiqueta"]: c["cantidad"] for c in datos["por_estado_aliado"]} == {
        "Aliado activo": 1,
        "Aliado inactivo": 1,
        "Contraparte propuesta (sin formalizar)": 1,
        "Sin aliado registrado": 1,
    }
    for activo in ("true", "false"):
        filtrado = client.get(URL_JSON, params={"aliado_activo": activo}).json()
        assert filtrado["total_convenios"] == 1


@pytest.mark.parametrize("con_pais_formal", [True, False])
def test_aliado_formal_tiene_prioridad_incluso_sin_pais(
    client, gestor, crear_convenio_estadistica, crear_aliado, crear_pais, con_pais_formal
) -> None:
    pais = crear_pais("Colombia") if con_pais_formal else None
    crear_convenio_estadistica(
        estado=EstadoConvenio.VIGENTE,
        aliado=crear_aliado(tipo=TipoAliado.EMPRESA, activo=False, pais=pais),
        nombre_propuesto="Universidad anterior",
        tipo_propuesto=TipoAliado.UNIVERSIDAD,
        pais_propuesto="México",
    )

    datos = client.get(URL_JSON).json()

    assert datos["por_tipo_aliado"] == [{"etiqueta": "Empresa", "cantidad": 1}]
    assert datos["por_estado_aliado"] == [{"etiqueta": "Aliado inactivo", "cantidad": 1}]
    assert datos["por_pais"] == [
        {"etiqueta": "Colombia" if con_pais_formal else "Sin país registrado", "cantidad": 1}
    ]
    for filtros in ({"pais": "México"}, {"tipo_aliado": "UNIVERSIDAD"}):
        assert client.get(URL_JSON, params=filtros).json()["total_convenios"] == 0


def test_convenios_compartidos_y_paises_normalizados_sin_n_mas_uno(
    db, client, gestor, crear_convenio_estadistica, crear_aliado, crear_pais
) -> None:
    aliado = crear_aliado(pais=crear_pais("Colombia"))
    for _ in range(3):
        crear_convenio_estadistica(estado=EstadoConvenio.VIGENTE, aliado=aliado)
    for pais in ("  COLOMBIA  ", "colombia"):
        crear_convenio_estadistica(
            estado=EstadoConvenio.EN_TRAMITE,
            tipo_propuesto=TipoAliado.UNIVERSIDAD,
            pais_propuesto=pais,
        )
    consultas = []

    def contar_select(_conn, _cursor, statement, _params, _context, _executemany):
        if statement.lstrip().upper().startswith("SELECT"):
            consultas.append(statement)

    event.listen(db.bind, "before_cursor_execute", contar_select)
    try:
        # Identity map nuevo para que los datos de las fixtures no oculten N+1.
        with Session(db.bind, join_transaction_mode="create_savepoint") as lectura:
            reporte = ServicioEstadisticasConvenios(lectura).generar()
    finally:
        event.remove(db.bind, "before_cursor_execute", contar_select)
    assert len(consultas) == 1
    assert reporte.total_convenios == 5
    assert reporte.por_pais == [Conteo(etiqueta="Colombia", cantidad=5)]
    assert reporte.por_tipo_aliado == [Conteo(etiqueta="Universidad", cantidad=5)]
    for campo in (
        "por_estado_contractual", "por_vigencia_activa", "por_estado_aliado",
        "por_tipo_aliado", "por_pais",
    ):
        assert sum(c.cantidad for c in getattr(reporte, campo)) == 5
    assert client.get(URL_JSON, params={"pais": " COLOMBIA "}).json()["total_convenios"] == 5


def test_filtros_propuestos_equivalentes_en_json_pdf_y_excel(
    client, gestor, crear_convenio_estadistica
) -> None:
    for pais in ("Costa Rica", " costa   RICA "):
        crear_convenio_estadistica(
            estado=EstadoConvenio.EN_TRAMITE,
            tipo_propuesto=TipoAliado.EMPRESA,
            pais_propuesto=pais,
        )
    crear_convenio_estadistica(
        estado=EstadoConvenio.FINALIZADO,
        tipo_propuesto=TipoAliado.EMPRESA,
        pais_propuesto="Costa Rica",
    )
    crear_convenio_estadistica(
        estado=EstadoConvenio.EN_TRAMITE,
        tipo_propuesto=TipoAliado.UNIVERSIDAD,
        pais_propuesto="Costa Rica",
    )
    filtros = {"estado": "EN_TRAMITE", "tipo_aliado": "EMPRESA", "pais": " COSTA  rica "}
    datos = client.get(URL_JSON, params=filtros).json()
    assert datos["total_convenios"] == 2
    excel = client.get(URL_EXCEL, params=filtros)
    assert excel.status_code == 200
    libro = load_workbook(BytesIO(excel.content))
    for hoja, campo in zip(libro, (
        "por_estado_contractual", "por_vigencia_activa", "por_estado_aliado",
        "por_tipo_aliado", "por_pais",
    ), strict=True):
        assert list(hoja.iter_rows(min_row=2, values_only=True)) == [
            (c["etiqueta"], c["cantidad"]) for c in datos[campo]
        ]
    pdf = client.get(URL_PDF, params=filtros)
    assert pdf.status_code == 200
    assert b"Costa Rica" in pdf.content
    assert b"Empresa" in pdf.content
    assert b"Contraparte propuesta" in pdf.content
    assert len(re.findall(rb"\(2\) Tj", pdf.content)) == 5


@pytest.mark.parametrize("prefijo", ["=", "+", "-", "@"])
def test_excel_fuerza_texto_literal_en_las_cinco_agrupaciones(prefijo) -> None:
    etiqueta = f"{prefijo}SUM(1,2)"
    conteos = [Conteo(etiqueta=etiqueta, cantidad=2)]
    reporte = ReporteEstadisticoConvenios(
        datos_suficientes=True, total_convenios=2,
        por_estado_contractual=conteos, por_vigencia_activa=conteos,
        por_estado_aliado=conteos, por_tipo_aliado=conteos, por_pais=conteos,
    )
    libro = load_workbook(BytesIO(generar_excel_estadisticas_convenios(reporte)))
    assert len(libro.sheetnames) == 5
    for hoja in libro:
        assert hoja["A2"].value == f"'{etiqueta}"
        assert hoja["A2"].data_type == "s"
        assert hoja["B2"].value == 2


@pytest.mark.parametrize("formal", [True, False])
@pytest.mark.parametrize("prefijo", ["=", "+", "-", "@"])
def test_excel_api_protege_paises_formales_y_propuestos(
    client, gestor, crear_convenio_estadistica, crear_aliado, crear_pais, formal, prefijo
) -> None:
    nombre = f"{prefijo}SUM(1,2)"
    aliado = crear_aliado(pais=crear_pais(nombre)) if formal else None
    crear_convenio_estadistica(
        estado=EstadoConvenio.VIGENTE, aliado=aliado, pais_propuesto=nombre
    )

    assert client.get(URL_JSON).json()["por_pais"] == [{"etiqueta": nombre, "cantidad": 1}]
    libro = load_workbook(BytesIO(client.get(URL_EXCEL).content))
    assert libro["País"]["A2"].value == f"'{nombre}"
    assert libro["País"]["A2"].data_type == "s"


def test_pdf_conserva_el_pais_como_texto_sin_interpretar_markup(
    client, gestor, crear_convenio_estadistica
) -> None:
    crear_convenio_estadistica(
        estado=EstadoConvenio.EN_TRAMITE, pais_propuesto="<b>Colombia</b> & Perú"
    )
    respuesta = client.get(URL_PDF)
    assert respuesta.status_code == 200
    # ReportLab escribe las entidades escapadas en varios operadores de texto.
    texto = b"".join(re.findall(rb"\((.*?)\) Tj", respuesta.content))
    assert b"<b>Colombia</b> &" in texto


def test_filtra_por_estado(client, gestor, crear_convenio_estadistica) -> None:
    crear_convenio_estadistica(estado=EstadoConvenio.VIGENTE)
    crear_convenio_estadistica(estado=EstadoConvenio.FINALIZADO)

    respuesta = client.get(URL_JSON, params={"estado": "VIGENTE"})

    assert respuesta.json()["total_convenios"] == 1


def test_filtra_por_estado_de_aliado(
    client, gestor, crear_convenio_estadistica, crear_aliado
) -> None:
    crear_convenio_estadistica(
        estado=EstadoConvenio.VIGENTE, aliado=crear_aliado(activo=True)
    )
    crear_convenio_estadistica(
        estado=EstadoConvenio.VIGENTE, aliado=crear_aliado(activo=False)
    )

    respuesta = client.get(URL_JSON, params={"aliado_activo": "true"})

    assert respuesta.json()["total_convenios"] == 1


def test_filtra_por_tipo_de_aliado(
    client, gestor, crear_convenio_estadistica, crear_aliado
) -> None:
    crear_convenio_estadistica(
        estado=EstadoConvenio.VIGENTE, aliado=crear_aliado(tipo=TipoAliado.UNIVERSIDAD)
    )
    crear_convenio_estadistica(
        estado=EstadoConvenio.VIGENTE, aliado=crear_aliado(tipo=TipoAliado.EMPRESA)
    )

    respuesta = client.get(URL_JSON, params={"tipo_aliado": "EMPRESA"})

    assert respuesta.json()["total_convenios"] == 1


def test_filtra_por_pais_sin_sensibilidad_a_mayusculas(
    client, gestor, crear_convenio_estadistica, crear_aliado, crear_pais
) -> None:
    colombia = crear_pais("Colombia")
    mexico = crear_pais("México")
    crear_convenio_estadistica(
        estado=EstadoConvenio.VIGENTE, aliado=crear_aliado(pais=colombia)
    )
    crear_convenio_estadistica(
        estado=EstadoConvenio.VIGENTE, aliado=crear_aliado(pais=mexico)
    )

    respuesta = client.get(URL_JSON, params={"pais": "colombia"})

    assert respuesta.json()["total_convenios"] == 1


def test_combina_varios_filtros(
    client, gestor, crear_convenio_estadistica, crear_aliado, crear_pais
) -> None:
    colombia = crear_pais("Colombia")
    coincide = crear_aliado(tipo=TipoAliado.EMPRESA, activo=True, pais=colombia)
    crear_convenio_estadistica(estado=EstadoConvenio.VIGENTE, aliado=coincide)
    # No coincide: mismo país y tipo, pero estado distinto.
    crear_convenio_estadistica(estado=EstadoConvenio.FINALIZADO, aliado=coincide)

    respuesta = client.get(
        URL_JSON,
        params={
            "estado": "VIGENTE",
            "tipo_aliado": "EMPRESA",
            "aliado_activo": "true",
            "pais": "Colombia",
        },
    )

    assert respuesta.json()["total_convenios"] == 1


# ---- CA-03: sin datos suficientes ----------------------------------------


def test_sin_convenios_reporta_datos_insuficientes(client, gestor) -> None:
    respuesta = client.get(URL_JSON)

    datos = respuesta.json()
    assert datos["datos_suficientes"] is False
    assert datos["total_convenios"] == 0
    assert datos["por_estado_contractual"] == []
    assert datos["por_vigencia_activa"] == []
    assert datos["por_estado_aliado"] == []
    assert datos["por_tipo_aliado"] == []
    assert datos["por_pais"] == []


def test_filtro_sin_coincidencias_tambien_reporta_datos_insuficientes(
    client, gestor, crear_convenio_estadistica
) -> None:
    crear_convenio_estadistica(estado=EstadoConvenio.VIGENTE)

    respuesta = client.get(URL_JSON, params={"estado": "CANCELADO"})

    assert respuesta.json()["datos_suficientes"] is False


def test_pdf_sin_datos_muestra_mensaje(client, gestor) -> None:
    respuesta = client.get(URL_PDF)

    assert respuesta.status_code == 200
    assert respuesta.content.startswith(b"%PDF")
    assert b"No hay datos suficientes" in respuesta.content


def test_excel_sin_datos_muestra_mensaje(client, gestor) -> None:
    respuesta = client.get(URL_EXCEL)

    assert respuesta.status_code == 200
    libro = load_workbook(BytesIO(respuesta.content))
    assert libro.sheetnames == ["Estadísticas"]
    assert "No hay datos suficientes" in libro["Estadísticas"]["A1"].value


# ---- Permisos ---------------------------------------------------------------


def test_requiere_autenticacion(client) -> None:
    assert client.get(URL_JSON).status_code == 401


def test_gestor_puede_consultar(client, gestor) -> None:
    assert client.get(URL_JSON).status_code == 200


def test_administrador_puede_consultar(client, crear_usuario, entrar_como) -> None:
    entrar_como(crear_usuario(CodigoRol.ADMINISTRADOR_ORI))
    assert client.get(URL_JSON).status_code == 200


@pytest.mark.parametrize(
    ("rol", "tipo"),
    [
        (CodigoRol.REVISOR_ORI, TipoUsuario.INTERNO),
        (CodigoRol.SOLICITANTE_INTERNO, TipoUsuario.INTERNO),
        (CodigoRol.SOLICITANTE_EXTERNO, TipoUsuario.EXTERNO),
    ],
)
def test_otros_roles_no_pueden_consultar(
    client, crear_usuario, entrar_como, rol, tipo
) -> None:
    entrar_como(crear_usuario(rol, tipo))
    assert client.get(URL_JSON).status_code == 403


def test_descarga_pdf_requiere_permiso(client, crear_usuario, entrar_como) -> None:
    entrar_como(crear_usuario(CodigoRol.REVISOR_ORI))
    assert client.get(URL_PDF).status_code == 403


def test_descarga_excel_requiere_permiso(client, crear_usuario, entrar_como) -> None:
    entrar_como(crear_usuario(CodigoRol.REVISOR_ORI))
    assert client.get(URL_EXCEL).status_code == 403


# ---- Descargas PDF / Excel con datos -----------------------------------------


def test_descarga_pdf_con_datos(
    client, gestor, crear_convenio_estadistica, crear_aliado, crear_pais
) -> None:
    colombia = crear_pais("Colombia")
    crear_convenio_estadistica(
        estado=EstadoConvenio.VIGENTE, aliado=crear_aliado(pais=colombia)
    )

    respuesta = client.get(URL_PDF)

    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"] == "application/pdf"
    assert "attachment" in respuesta.headers["content-disposition"]
    assert respuesta.content.startswith(b"%PDF")
    assert b"Colombia" in respuesta.content


def test_descarga_excel_con_datos(
    client, gestor, crear_convenio_estadistica, crear_aliado, crear_pais
) -> None:
    colombia = crear_pais("Colombia")
    crear_convenio_estadistica(
        estado=EstadoConvenio.VIGENTE, aliado=crear_aliado(pais=colombia)
    )
    crear_convenio_estadistica(estado=EstadoConvenio.FINALIZADO, aliado=None)

    respuesta = client.get(URL_EXCEL)

    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in respuesta.headers["content-disposition"]

    libro = load_workbook(BytesIO(respuesta.content))
    assert libro.sheetnames == [
        "Estado contractual",
        "Vigencia activa",
        "Estado del aliado",
        "Tipo de aliado",
        "País",
    ]
    hoja_pais = libro["País"]
    filas = [tuple(celda.value for celda in fila) for fila in hoja_pais.iter_rows(min_row=2)]
    assert ("Colombia", 1) in filas

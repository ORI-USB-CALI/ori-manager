from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.exc import SQLAlchemyError

from backend.models.borrador_convenio import BorradorConvenio
from backend.models.convenio import Convenio
from backend.models.enums import ContextoVersionConvenio
from backend.models.revision_convenio import RevisionConvenio
from backend.models.version_convenio import VersionConvenio
from backend.schemas.convenio import ConvenioElaboracionGuardar
from backend.services.convenios import ConflictoBorradorConvenio, ServicioConvenios
from backend.services.notificaciones import ServicioNotificaciones


def _ruta(convenio_id: int) -> str:
    return f"/api/convenios/{convenio_id}/elaboracion"


def _elaboracion(client, convenio_id: int) -> dict:
    respuesta = client.get(_ruta(convenio_id))
    assert respuesta.status_code == 200
    return respuesta.json()


def _con_parrafo(contenido: dict, texto: str) -> dict:
    nuevo = deepcopy(contenido)
    nuevo["content"].append(
        {"type": "paragraph", "content": [{"type": "text", "text": texto}]}
    )
    return nuevo


def _guardar(client, convenio_id: int, contenido: dict, revision: int, **datos):
    return client.patch(
        _ruta(convenio_id),
        json={"contenido": contenido, "expected_revision": revision, **datos},
    )


def _versiones(db, convenio_id: int) -> list[VersionConvenio]:
    db.expire_all()
    return list(
        db.scalars(
            select(VersionConvenio)
            .where(VersionConvenio.convenio_id == convenio_id)
            .order_by(VersionConvenio.numero)
        )
    )


def _revision_pendiente(db, convenio_id: int) -> RevisionConvenio:
    return db.scalar(
        select(RevisionConvenio).where(
            RevisionConvenio.convenio_id == convenio_id,
            RevisionConvenio.estado == "PENDIENTE",
        )
    )


def _texto(nodo: dict) -> str:
    partes = [str(nodo.get("text", ""))]
    partes.extend(_texto(hijo) for hijo in nodo.get("content", []))
    return " ".join(partes)


# ---- CA-01 y CA-02: guardar avances y recuperarlos ----------------------------


def test_ca01_guardar_avance_persiste_contenido_y_datos_sin_versionar(
    client, db, convenio_listo
) -> None:
    inicial = _elaboracion(client, convenio_listo.id)
    contenido = _con_parrafo(inicial["contenido"], "Avance del Gestor")

    respuesta = _guardar(
        client, convenio_listo.id, contenido, 0, objeto="Objeto ajustado en avance"
    )

    assert respuesta.status_code == 200, respuesta.text
    cuerpo = respuesta.json()
    assert cuerpo["version_actual"] == inicial["version_actual"]
    assert cuerpo["revision_borrador"] == 1
    assert cuerpo["borrador_guardado_en"] is not None
    assert cuerpo["contenido"] == contenido
    assert cuerpo["objeto"] == "Objeto ajustado en avance"
    assert [v.numero for v in _versiones(db, convenio_listo.id)] == [1]


def test_ca02_el_avance_se_recupera_y_las_versiones_no_cambian(
    client, db, convenio_listo
) -> None:
    original = _versiones(db, convenio_listo.id)[0].contenido
    contenido = _con_parrafo(original, "Trabajo de ayer")
    assert _guardar(client, convenio_listo.id, contenido, 0).status_code == 200
    segundo = _con_parrafo(contenido, "Trabajo de hoy")
    assert _guardar(client, convenio_listo.id, segundo, 1).status_code == 200

    recuperado = _elaboracion(client, convenio_listo.id)

    assert recuperado["contenido"] == segundo
    assert recuperado["revision_borrador"] == 2
    (version,) = _versiones(db, convenio_listo.id)
    assert version.contenido == original


# ---- CA-03: crear o reutilizar la versión al enviar a revisión ----------------


def test_ca03_enviar_crea_la_version_con_el_avance_y_vincula_la_revision(
    client, db, convenio_listo
) -> None:
    contenido = _con_parrafo(
        _elaboracion(client, convenio_listo.id)["contenido"], "Lista para Jurídica"
    )
    guardado = _guardar(client, convenio_listo.id, contenido, 0).json()

    enviado = client.post(
        f"{_ruta(convenio_listo.id)}/finalizar",
        json={"expected_revision": guardado["revision_borrador"]},
    )

    assert enviado.status_code == 200, enviado.text
    versiones = _versiones(db, convenio_listo.id)
    assert [v.numero for v in versiones] == [1, 2]
    assert versiones[1].contenido == contenido
    assert versiones[1].contexto == ContextoVersionConvenio.FINALIZACION.value
    assert _revision_pendiente(db, convenio_listo.id).version_convenio_id == (
        versiones[1].id
    )
    assert db.scalar(
        select(BorradorConvenio).where(
            BorradorConvenio.convenio_id == convenio_listo.id
        )
    ) is None


def test_ca03_enviar_sin_cambios_reutiliza_la_version_sin_duplicarla(
    client, db, convenio_listo
) -> None:
    (v1,) = _versiones(db, convenio_listo.id)
    assert _guardar(client, convenio_listo.id, v1.contenido, 0).status_code == 200

    enviado = client.post(f"{_ruta(convenio_listo.id)}/finalizar")

    assert enviado.status_code == 200, enviado.text
    assert [v.numero for v in _versiones(db, convenio_listo.id)] == [1]
    assert _revision_pendiente(db, convenio_listo.id).version_convenio_id == v1.id


def test_ca03_un_cambio_solo_de_datos_tambien_genera_version(
    client, db, convenio_listo
) -> None:
    (v1,) = _versiones(db, convenio_listo.id)
    guardado = _guardar(
        client,
        convenio_listo.id,
        v1.contenido,
        0,
        implicacion_financiera="Aporte en especie de ambas partes",
    )
    assert guardado.status_code == 200, guardado.text

    assert client.post(f"{_ruta(convenio_listo.id)}/finalizar").status_code == 200

    versiones = _versiones(db, convenio_listo.id)
    assert [v.numero for v in versiones] == [1, 2]
    assert versiones[1].snapshot_metadata["implicacion_financiera"] == (
        "Aporte en especie de ambas partes"
    )


# ---- CA-04 y CA-05: ciclos de corrección sin tocar versiones anteriores --------


def test_ca04_correcciones_parciales_no_versionan_y_el_reenvio_si(
    client, db, gestor, revisor, convenio_listo
) -> None:
    servicio = ServicioConvenios(db)
    servicio.finalizar_elaboracion(convenio_listo.id, gestor)
    revision_v1 = _revision_pendiente(db, convenio_listo.id)
    v1 = db.get(VersionConvenio, revision_v1.version_convenio_id)
    contenido_v1 = deepcopy(v1.contenido)
    devuelta = servicio.devolver(
        convenio_listo.id,
        revision_v1.id,
        ["Precisar el objeto del convenio"],
        convenio_listo.version_actual,
        revisor,
    )
    servicio.atender_observacion(
        convenio_listo.id, devuelta.observaciones[0].id, "Objeto precisado", gestor
    )

    elaboracion = _elaboracion(client, convenio_listo.id)
    assert elaboracion["contenido"] == contenido_v1
    assert elaboracion["revision_borrador"] == 0
    parcial = _con_parrafo(contenido_v1, "Corrección parcial")
    assert _guardar(client, convenio_listo.id, parcial, 0).status_code == 200
    final = _con_parrafo(parcial, "Corrección final")
    assert _guardar(client, convenio_listo.id, final, 1).status_code == 200
    assert [v.numero for v in _versiones(db, convenio_listo.id)] == [v1.numero]

    reenvio = client.post(
        f"{_ruta(convenio_listo.id)}/finalizar", json={"expected_revision": 2}
    )

    assert reenvio.status_code == 200, reenvio.text
    versiones = _versiones(db, convenio_listo.id)
    assert [v.numero for v in versiones] == [v1.numero, v1.numero + 1]
    assert versiones[0].contenido == contenido_v1
    assert versiones[1].contenido == final
    assert db.get(RevisionConvenio, revision_v1.id).version_convenio_id == v1.id
    assert _revision_pendiente(db, convenio_listo.id).version_convenio_id == (
        versiones[1].id
    )


# ---- CA-06: concurrencia del borrador ----------------------------------------


def test_ca06_enviar_con_una_revision_vieja_responde_409_sin_cambios(
    client, db, convenio_listo
) -> None:
    contenido = _con_parrafo(
        _elaboracion(client, convenio_listo.id)["contenido"], "Guardado por otra persona"
    )
    assert _guardar(client, convenio_listo.id, contenido, 0).status_code == 200

    conflicto = client.post(
        f"{_ruta(convenio_listo.id)}/finalizar", json={"expected_revision": 0}
    )

    assert conflicto.status_code == 409
    assert conflicto.json()["detail"]["current_revision"] == 1
    db.expire_all()
    assert db.get(Convenio, convenio_listo.id).etapa_actual.codigo == "ELABORACION"
    assert [v.numero for v in _versiones(db, convenio_listo.id)] == [1]


# ---- CA-07: integridad de la transición --------------------------------------


def test_ca07_un_error_al_abrir_la_revision_no_deja_cambios_parciales(
    client, db, gestor, convenio_listo, monkeypatch
) -> None:
    contenido = _con_parrafo(
        _elaboracion(client, convenio_listo.id)["contenido"], "Envío que falla"
    )
    assert _guardar(client, convenio_listo.id, contenido, 0).status_code == 200

    def fallar(*_args, **_kwargs):
        raise SQLAlchemyError("falla simulada al notificar")

    monkeypatch.setattr(ServicioNotificaciones, "crear_para_rol", fallar)
    with pytest.raises(SQLAlchemyError):
        ServicioConvenios(db).finalizar_elaboracion(convenio_listo.id, gestor)

    db.expire_all()
    assert db.get(Convenio, convenio_listo.id).etapa_actual.codigo == "ELABORACION"
    assert [v.numero for v in _versiones(db, convenio_listo.id)] == [1]
    assert db.scalar(
        select(func.count())
        .select_from(RevisionConvenio)
        .where(RevisionConvenio.convenio_id == convenio_listo.id)
    ) == 0
    borrador = db.scalar(
        select(BorradorConvenio).where(
            BorradorConvenio.convenio_id == convenio_listo.id
        )
    )
    assert (borrador.revision, borrador.contenido) == (1, contenido)


def test_ca07_guardar_sin_cambios_no_sube_la_revision(
    client, db, convenio_listo
) -> None:
    contenido = _elaboracion(client, convenio_listo.id)["contenido"]

    respuesta = _guardar(client, convenio_listo.id, contenido, 0)

    assert respuesta.status_code == 200
    assert respuesta.json()["revision_borrador"] == 0
    assert db.scalar(
        select(BorradorConvenio).where(
            BorradorConvenio.convenio_id == convenio_listo.id
        )
    ) is None


def test_ca07_el_conflicto_desde_el_servicio_conserva_el_avance(
    client, db, gestor, convenio_listo
) -> None:
    contenido = _con_parrafo(
        _elaboracion(client, convenio_listo.id)["contenido"], "Primer guardado"
    )
    assert _guardar(client, convenio_listo.id, contenido, 0).status_code == 200

    with pytest.raises(ConflictoBorradorConvenio):
        ServicioConvenios(db).guardar_elaboracion(
            convenio_listo.id,
            ConvenioElaboracionGuardar(
                contenido=_con_parrafo(contenido, "Pisaría"), expected_revision=0
            ),
            gestor,
        )

    assert _elaboracion(client, convenio_listo.id)["contenido"] == contenido


# ---- CA-08: compatibilidad con convenios anteriores --------------------------


def test_ca08_un_convenio_sin_borrador_trabaja_sobre_su_version_actual(
    client, db, gestor, convenio_listo
) -> None:
    (v1,) = _versiones(db, convenio_listo.id)
    historica = VersionConvenio(
        convenio_id=convenio_listo.id,
        numero=2,
        contenido=_con_parrafo(v1.contenido, "Guardado anterior a HT-13"),
        snapshot_metadata=v1.snapshot_metadata,
        autor_id=gestor.id,
        etapa_id=v1.etapa_id,
        contexto=ContextoVersionConvenio.GUARDADO.value,
    )
    db.add(historica)
    convenio_listo.version_actual = 2
    db.commit()

    elaboracion = _elaboracion(client, convenio_listo.id)

    assert elaboracion["version_actual"] == 2
    assert elaboracion["revision_borrador"] == 0
    assert "Guardado anterior a HT-13" in _texto(elaboracion["contenido"])
    assert client.post(f"{_ruta(convenio_listo.id)}/finalizar").status_code == 200
    assert [v.numero for v in _versiones(db, convenio_listo.id)] == [1, 2]


# ---- Exportación del documento de trabajo ------------------------------------


def test_en_elaboracion_se_exporta_el_avance_guardado(
    client, db, convenio_listo
) -> None:
    contenido = _con_parrafo(
        _elaboracion(client, convenio_listo.id)["contenido"], "Cláusula del avance"
    )
    assert _guardar(client, convenio_listo.id, contenido, 0).status_code == 200

    respuesta = client.get(
        f"/api/convenios/{convenio_listo.id}/exportar", params={"formato": "excel"}
    )

    assert respuesta.status_code == 200, respuesta.text
    libro = load_workbook(BytesIO(respuesta.content))
    textos = {
        str(celda.value)
        for hoja in libro.worksheets
        for fila in hoja.iter_rows()
        for celda in fila
        if celda.value is not None
    }
    assert any("Cláusula del avance" in texto for texto in textos)

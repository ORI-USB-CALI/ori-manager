"""HU-24: correcciones de revisión del PR #120.

Autorización por alcance, rutas oficiales, aislamiento de datos, versión
contractual, texto especial, contenido incompleto y exportación segura.
"""
import re
from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from backend.core.roles import CodigoRol, TipoUsuario
from backend.models.enums import ContextoVersionConvenio
from backend.models.etapa import Etapa
from backend.models.firma_convenio import FirmaConvenio
from backend.models.proceso_firmas_convenio import ProcesoFirmasConvenio
from backend.models.revision_convenio import RevisionConvenio
from backend.models.version_convenio import VersionConvenio
from backend.schemas.exportacion import (
    INFORMACION_NO_DISPONIBLE,
    ConvenioPdfDTO,
    FirmantePdfDTO,
    SeccionContenidoDTO,
)
from backend.services.convenios import ServicioConvenios
from backend.services.excel_convenio import (
    HOJA_DATOS_GENERALES,
    HOJA_PARTICIPANTES_COMPROMISOS,
    ConvenioExcelExportService,
)
from backend.services.pdf_convenio import ServicioPdfConvenio

TEXTO_ESPECIAL = 'Acuerdo <b>sin cerrar & "comillas" <script>alert(1)</script> \' > <'


def _exportar(client, convenio_id, formato="pdf"):
    return client.get(f"/api/convenios/{convenio_id}/exportar?formato={formato}")


def _dto(**cambios) -> ConvenioPdfDTO:
    datos = {
        "convenio_id": 1,
        "modo_emision": "OFICIAL",
        "codigo": "CONV-0001",
        "solicitud_consecutivo": "SOL-0001",
        "tipo_convenio": "Marco",
        "estado": "EN_TRAMITE",
        "etapa": "Elaboración",
        "alcance": "INSTITUCIONAL",
        "objeto": "Cooperación académica",
        "justificacion": "Fortalecer la movilidad",
        "implicacion_financiera": "Sin costo",
        "fecha_inicio": date(2026, 1, 1),
        "fecha_vencimiento": date(2028, 1, 1),
        "duracion_meses": 24,
        "porcentaje_avance": None,
        "version_actual": 1,
        "solicitante_nombre": "Ana Gestora",
        "solicitante_correo": "ana@usbcali.edu.co",
        "solicitante_unidad": "Facultad de Ingeniería",
        "solicitante_tipo": "INTERNO",
        "aliado_nombre": "Universidad Aliada",
        "aliado_tipo": "UNIVERSIDAD",
        "aliado_identificacion": "900123456",
        "aliado_pais": "Perú",
        "aliado_ciudad": "Lima",
        "aliado_contacto_nombre": "Luis Contacto",
        "aliado_contacto_correo": "luis@aliada.example",
        "secciones_contenido": [],
        "firmantes": [],
    }
    datos.update(cambios)
    return ConvenioPdfDTO(**datos)


def _textos_libro(contenido: bytes) -> set[str]:
    libro = load_workbook(BytesIO(contenido))
    return {
        str(celda.value)
        for hoja in libro.worksheets
        for fila in hoja.iter_rows()
        for celda in fila
        if celda.value is not None
    }


# ---- Autorización por alcance --------------------------------------------------


@pytest.mark.parametrize("formato", ["pdf", "excel"])
def test_revisor_ori_no_exporta_un_convenio_ajeno_a_su_alcance(
    client, crear_usuario, entrar_como, convenio_exportable, formato
):
    # El convenio está completo pero no en revisión jurídica: fuera del alcance
    # del Revisor ORI aunque tenga el permiso convenios.ver.
    entrar_como(crear_usuario(CodigoRol.REVISOR_ORI, TipoUsuario.INTERNO))

    respuesta = _exportar(client, convenio_exportable.id, formato)

    assert respuesta.status_code == 404
    # No revela el contenido ni qué datos le faltan al convenio.
    assert "campos_faltantes" not in respuesta.text
    assert convenio_exportable.codigo not in respuesta.text


def test_revisor_ori_con_revision_juridica_elegible_supera_el_alcance(
    db, client, gestor, crear_usuario, entrar_como, convenio_listo
):
    ServicioConvenios(db).finalizar_elaboracion(convenio_listo.id, gestor)
    entrar_como(crear_usuario(CodigoRol.REVISOR_ORI, TipoUsuario.INTERNO))

    respuesta = _exportar(client, convenio_listo.id)

    # Dentro de su alcance llega a la validación (el convenio está incompleto).
    assert respuesta.status_code == 422
    assert respuesta.json()["error"] == "CONVENIO_INCOMPLETO"


def test_administrador_ori_exporta_cualquier_convenio(
    client, crear_usuario, entrar_como, convenio_exportable
):
    entrar_como(crear_usuario(CodigoRol.ADMINISTRADOR_ORI, TipoUsuario.INTERNO))

    assert _exportar(client, convenio_exportable.id).status_code == 200


def test_las_rutas_de_convenios_ya_no_se_publican_bajo_api_v1(
    client, gestor, convenio_exportable
):
    assert client.get(f"/api/v1/convenios/{convenio_exportable.id}").status_code == 404
    assert client.get(f"/api/convenios/{convenio_exportable.id}").status_code == 200


# ---- Aislamiento de datos ------------------------------------------------------


def test_la_exportacion_no_mezcla_datos_de_otro_convenio(
    db, client, gestor, crear_convenio, convenio_exportable
):
    otro = crear_convenio(gestor, objeto="Objeto exclusivo del otro convenio")
    otro.codigo = "CONV-AJENO-0001"
    db.commit()

    respuesta = _exportar(client, convenio_exportable.id, "excel")

    assert respuesta.status_code == 200
    textos = _textos_libro(respuesta.content)
    assert convenio_exportable.codigo in textos
    assert "CONV-AJENO-0001" not in textos
    assert "Objeto exclusivo del otro convenio" not in textos


# ---- Versión contractual -------------------------------------------------------


def test_version_indicada_inexistente_responde_422_sin_usar_otra(
    db, client, gestor, convenio_exportable
):
    convenio_exportable.version_actual = 99
    db.commit()

    respuesta = _exportar(client, convenio_exportable.id)

    assert respuesta.status_code == 422
    assert "version_contractual" in respuesta.json()["campos_faltantes"]


def _mover_a_etapa(db, convenio, codigo: str) -> None:
    convenio.etapa_actual = db.scalar(select(Etapa).where(Etapa.codigo == codigo))
    db.commit()


def _crear_version_2(client, db, convenio, texto: str) -> None:
    """Crea la versión 2 como lo hace el envío a revisión: desde HT-13, guardar
    avance ya no versiona. La etapa no cambia, así que la exportación sigue
    siendo oficial."""
    contenido = {
        "type": "doc",
        "content": [{"type": "paragraph", "content": [{"type": "text", "text": texto}]}],
    }
    ServicioConvenios(db)._crear_version(
        convenio, contenido, convenio.creado_por, ContextoVersionConvenio.FINALIZACION
    )
    db.commit()


def _completar_proceso_de_la_v1(db, convenio) -> None:
    proceso = db.scalar(
        select(ProcesoFirmasConvenio).where(
            ProcesoFirmasConvenio.convenio_id == convenio.id
        )
    )
    proceso.estado = "COMPLETADO"
    db.commit()


def _proceso_de_la_version_actual(db, convenio, creado_por, rector, representante):
    """Proceso de firmas propio de la versión actual, como en convenio_exportable."""
    db.expire_all()
    version = db.scalar(
        select(VersionConvenio).where(
            VersionConvenio.convenio_id == convenio.id,
            VersionConvenio.numero == convenio.version_actual,
        )
    )
    revision_final = RevisionConvenio(
        convenio_id=convenio.id,
        tipo="FINAL",
        version_convenio_id=version.id,
        estado="RESUELTA",
        resultado="APROBADA",
    )
    db.add(revision_final)
    db.flush()
    proceso = ProcesoFirmasConvenio(
        convenio_id=convenio.id,
        version_convenio_id=version.id,
        revision_final_id=revision_final.id,
        creado_por_id=creado_por.id,
        estado="CONFIGURACION",
    )
    db.add(proceso)
    db.flush()
    db.add_all(
        [
            FirmaConvenio(
                proceso_firmas_id=proceso.id,
                orden=6,
                rol_firmante="RECTOR",
                parte="UNIVERSIDAD",
                nombre_firmante=rector,
                cargo_firmante="Rector",
                estado="PENDIENTE",
            ),
            FirmaConvenio(
                proceso_firmas_id=proceso.id,
                orden=7,
                rol_firmante="PARTE_SOLICITANTE",
                parte="REPRESENTANTE_LEGAL_ENTIDAD",
                nombre_firmante=representante,
                cargo_firmante="Representante legal",
                estado="PENDIENTE",
            ),
        ]
    )
    db.commit()


def test_v2_sin_proceso_rechaza_aunque_la_v1_tenga_proceso_completado(
    db, client, gestor, convenio_exportable
):
    _completar_proceso_de_la_v1(db, convenio_exportable)
    _crear_version_2(client, db, convenio_exportable, "Texto de la v2")

    respuesta = _exportar(client, convenio_exportable.id)

    # Rechaza por los requisitos de la versión actual, no por el histórico.
    assert respuesta.status_code == 422
    campos = respuesta.json()["campos_faltantes"]
    assert campos == ["representante_legal_institucion", "representante_legal_entidad"]


@pytest.mark.parametrize("estado_v1", ["COMPLETADO", "CANCELADO"])
def test_v2_con_su_proceso_se_exporta_solo_con_sus_firmantes(
    db, client, gestor, convenio_exportable, estado_v1
):
    proceso_v1 = db.scalar(
        select(ProcesoFirmasConvenio).where(
            ProcesoFirmasConvenio.convenio_id == convenio_exportable.id
        )
    )
    proceso_v1.estado = estado_v1
    db.commit()
    _crear_version_2(client, db, convenio_exportable, "Texto de la v2")
    _proceso_de_la_version_actual(
        db, convenio_exportable, gestor, "Rector de la v2", "Representante de la v2"
    )

    respuesta = _exportar(client, convenio_exportable.id, "excel")

    assert respuesta.status_code == 200
    textos = _textos_libro(respuesta.content)
    assert {"Rector de la v2", "Representante de la v2", "Texto de la v2"} <= textos
    # Los firmantes de la v1 (del fixture) nunca aparecen en la versión actual.
    assert "Rector de Prueba" not in textos
    assert "Representante Legal de Prueba" not in textos
    dto = ServicioPdfConvenio().mapear_convenio_a_dto(convenio_exportable.id, db)
    assert [firmante.nombre for firmante in dto.firmantes] == [
        "Rector de la v2",
        "Representante de la v2",
    ]


def test_el_mapeo_usa_el_contenido_de_la_version_actual(
    db, client, gestor, convenio_exportable
):
    _crear_version_2(client, db, convenio_exportable, "Solo en la v2")

    dto = ServicioPdfConvenio().mapear_convenio_a_dto(convenio_exportable.id, db)

    assert [seccion.texto for seccion in dto.secciones_contenido] == ["Solo en la v2"]
    # Los firmantes del proceso de la v1 no se atribuyen a la v2.
    assert dto.firmantes == []


# ---- Valores ausentes: nunca supuestos -----------------------------------------


def test_el_mapeo_no_inventa_ni_sustituye_valores_ausentes(
    db, gestor, convenio_exportable
):
    convenio_exportable.implicacion_financiera = None
    convenio_exportable.objeto = None
    convenio_exportable.aliado.ciudad = None
    convenio_exportable.aliado.pais_id = None
    solicitud = convenio_exportable.solicitud
    # Datos de la solicitud que antes se usaban como sustitutos silenciosos.
    solicitud.implicacion_financiera = "Valor propuesto en la solicitud"
    solicitud.pais_aliado_propuesto = "País propuesto"
    solicitud.ciudad_aliado_propuesto = "Ciudad propuesta"
    db.commit()
    db.expire_all()

    dto = ServicioPdfConvenio().mapear_convenio_a_dto(convenio_exportable.id, db)

    assert dto.implicacion_financiera == INFORMACION_NO_DISPONIBLE
    assert dto.objeto == INFORMACION_NO_DISPONIBLE
    assert dto.aliado_pais == INFORMACION_NO_DISPONIBLE
    assert dto.aliado_ciudad == INFORMACION_NO_DISPONIBLE
    for valor in ("Sin costo para la Universidad", "Colombia", "Cali"):
        assert valor not in dto.model_dump_json()


# ---- Texto especial y formato del documento estructurado -----------------------


def test_texto_de_usuario_no_se_interpreta_como_marcado_en_el_pdf():
    dto = _dto(
        objeto=TEXTO_ESPECIAL,
        justificacion=TEXTO_ESPECIAL,
        aliado_nombre="Aliado <i>& Cía",
        implicacion_financiera="<para>sin cerrar",
        firmantes=[
            FirmantePdfDTO(
                orden=6,
                rol="Rectoría",
                parte="UNIVERSIDAD",
                nombre="Nombre <b>",
                cargo="Cargo & <br>",
                modalidad="FISICA",
                estado="PENDIENTE",
            )
        ],
    )

    # Sin escape, ReportLab rechaza el marcado mal formado al renderizar.
    contenido = ServicioPdfConvenio().renderizar_dto_a_pdf(dto)

    assert contenido.startswith(b"%PDF-")


def test_las_secciones_escapan_el_texto_y_conservan_los_formatos_validos():
    documento = {
        "type": "doc",
        "content": [
            {
                "type": "heading",
                "attrs": {"level": 2},
                "content": [{"type": "text", "text": "Cláusula <1> & partes"}],
            },
            {
                "type": "paragraph",
                "content": [
                    {"type": "text", "text": "<b>no</b>", "marks": [{"type": "bold"}]},
                    {"type": "hardBreak"},
                    {"type": "text", "text": "cursiva", "marks": [{"type": "italic"}]},
                ],
            },
            {
                "type": "orderedList",
                "attrs": {"start": 3},
                "content": [
                    {
                        "type": "listItem",
                        "content": [
                            {"type": "paragraph", "content": [{"type": "text", "text": "Tercera"}]},
                            {
                                "type": "bulletList",
                                "content": [
                                    {
                                        "type": "listItem",
                                        "content": [
                                            {
                                                "type": "paragraph",
                                                "content": [{"type": "text", "text": "Anidada"}],
                                            }
                                        ],
                                    }
                                ],
                            },
                        ],
                    }
                ],
            },
        ],
    }

    secciones = ServicioPdfConvenio()._extraer_secciones_contenido(documento)

    titulo, parrafo, item, anidado = secciones
    assert titulo.tipo == "heading" and titulo.nivel == 2
    assert titulo.texto == "Cláusula <1> & partes"
    assert titulo.texto_formateado == "Cláusula &lt;1&gt; &amp; partes"
    # La marca bold es marcado válido; las etiquetas escritas por el usuario no.
    assert parrafo.texto == "<b>no</b>\ncursiva"
    assert parrafo.texto_formateado == "<b>&lt;b&gt;no&lt;/b&gt;</b><br/><i>cursiva</i>"
    # La lista numerada respeta su número de inicio y la anidada su nivel.
    assert item.texto == "3. Tercera"
    assert anidado.texto == "  • Anidada"
    assert anidado.texto_formateado.startswith("&nbsp;")


def test_pdf_renderiza_las_secciones_formateadas():
    dto = _dto(
        secciones_contenido=[
            SeccionContenidoDTO(
                tipo="paragraph",
                texto="<b>literal</b>",
                texto_formateado="<b>&lt;b&gt;literal&lt;/b&gt;</b>",
            )
        ]
    )

    assert ServicioPdfConvenio().renderizar_dto_a_pdf(dto).startswith(b"%PDF-")


# ---- Contenido incompleto ------------------------------------------------------


def test_pdf_sin_contenido_no_redacta_clausulas_supletorias():
    servicio = ServicioPdfConvenio()
    elementos: list = []

    servicio._construir_clausulas_contractuales(
        _dto(secciones_contenido=[]), servicio._configurar_estilos(), elementos
    )

    textos = [getattr(elemento, "text", "") for elemento in elementos]
    assert "Contenido contractual no disponible." in textos
    assert not any("CLÁUSULA" in texto for texto in textos)


# ---- Exportación segura --------------------------------------------------------


@pytest.mark.parametrize(
    "valor", ['=HYPERLINK("http://malicioso")', "+cmd|' /C calc'!A0", "-2+3", "@SUM(A1:A2)"]
)
def test_excel_no_permite_formulas(valor):
    dto = _dto(objeto=valor, aliado_nombre=valor)

    libro = load_workbook(BytesIO(ConvenioExcelExportService().construir_libro(dto)))

    celdas = [
        celda
        for hoja in libro.worksheets
        for fila in hoja.iter_rows()
        for celda in fila
        if isinstance(celda.value, str) and valor in celda.value
    ]
    assert celdas
    assert all(celda.value == f"'{valor}" and celda.data_type == "s" for celda in celdas)


def test_excel_usa_el_texto_plano_de_las_clausulas():
    dto = _dto(
        secciones_contenido=[
            SeccionContenidoDTO(
                tipo="paragraph",
                texto="Primera: objeto",
                texto_formateado="<b>Primera:</b> objeto",
            )
        ]
    )

    libro = load_workbook(BytesIO(ConvenioExcelExportService().construir_libro(dto)))

    textos = _textos_libro(ConvenioExcelExportService().construir_libro(dto))
    assert "Primera: objeto" in textos
    assert not any("<b>" in texto for texto in textos)
    assert libro.sheetnames == [HOJA_DATOS_GENERALES, HOJA_PARTICIPANTES_COMPROMISOS]


def test_nombre_del_archivo_solo_usa_caracteres_seguros(
    db, client, gestor, convenio_exportable
):
    convenio_exportable.codigo = "CONV/../ñ✓ 01"
    db.commit()

    respuesta = _exportar(client, convenio_exportable.id)

    assert respuesta.status_code == 200
    assert re.fullmatch(
        r'attachment; filename="Convenio_[A-Za-z0-9_.-]+\.pdf"',
        respuesta.headers["content-disposition"],
    )
    assert "/" not in respuesta.headers["content-disposition"].split("filename=")[1]

"""HU-10/4A: exportación real de ambos orígenes, sin solicitudes ficticias."""

import re
import zlib
from base64 import a85decode
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from backend.models.aliado import Aliado
from backend.models.auditoria import Auditoria
from backend.models.convenio import Convenio
from backend.models.etapa import Etapa
from backend.models.firma_convenio import FirmaConvenio
from backend.models.historial_etapa import HistorialEtapa
from backend.models.proceso_firmas_convenio import ProcesoFirmasConvenio
from backend.models.revision_convenio import RevisionConvenio
from backend.models.solicitud_convenio import SolicitudConvenio
from backend.models.tipo_convenio import TipoConvenio
from backend.models.version_convenio import VersionConvenio
from backend.services.pdf_convenio import ServicioPdfConvenio
from backend.services.validador_exportacion_convenio import ConvenioExportValidator

NO_APLICA = "No aplica (creación directa ORI)"
NO_DISPONIBLE = "Información no disponible"


@pytest.fixture
def manual_documentado(client, db, gestor):
    aliado = Aliado(
        nombre="Universidad aliada del convenio manual",
        tipo="UNIVERSIDAD",
        tipo_identificacion="NIT",
        identificacion=f"9{uuid4().int % 10**8:08d}",
        correo="convenios-manual@example.com",
        ciudad="Bogotá",
    )
    db.add(aliado)
    db.flush()
    respuesta = client.post("/api/convenios/manual", json={
        "aliado_id": aliado.id,
        "codigo": f"MAN-{uuid4().hex[:8].upper()}",
        "objeto": "Cooperación académica desde la ORI",
        "tipo_convenio_id": db.scalar(select(TipoConvenio.id).where(TipoConvenio.codigo == "MARCO")),
        "alcance": "INSTITUCIONAL",
        "implicacion_financiera": "Sin compromiso financiero",
        "duracion_meses": 24,
        "fecha_inicio": "2026-01-01",
        "fecha_vencimiento": "2028-01-01",
    })
    assert respuesta.status_code == 201, respuesta.text
    convenio_id = respuesta.json()["id"]
    guardado = client.patch(f"/api/convenios/{convenio_id}/elaboracion", json={
        "expected_version": respuesta.json()["version_actual"],
        "contenido": {"type": "doc", "content": [{"type": "paragraph", "content": [
            {"type": "text", "text": "Cláusula documental de auditoría"},
        ]}]},
    })
    assert guardado.status_code == 200, guardado.text
    db.expire_all()
    return db.get(Convenio, convenio_id)


@pytest.fixture
def manual_formalizado(db, gestor, manual_documentado):
    """Estado persistido de exportación oficial, sin crear una solicitud."""
    convenio = manual_documentado
    version = db.scalar(select(VersionConvenio).where(
        VersionConvenio.convenio_id == convenio.id,
        VersionConvenio.numero == convenio.version_actual,
    ))
    revision = RevisionConvenio(
        convenio_id=convenio.id, tipo="FINAL", version_convenio_id=version.id,
        estado="RESUELTA", resultado="APROBADA",
    )
    db.add(revision)
    db.flush()
    proceso = ProcesoFirmasConvenio(
        convenio_id=convenio.id, version_convenio_id=version.id,
        revision_final_id=revision.id, creado_por_id=gestor.id, estado="COMPLETADO",
    )
    db.add(proceso)
    db.flush()
    db.add_all([
        FirmaConvenio(
            proceso_firmas_id=proceso.id, orden=6, rol_firmante="RECTOR",
            parte="UNIVERSIDAD", nombre_firmante="Rector de prueba",
            cargo_firmante="Rector", estado="FIRMADA",
        ),
        FirmaConvenio(
            proceso_firmas_id=proceso.id, orden=7, rol_firmante="PARTE_SOLICITANTE",
            parte="REPRESENTANTE_LEGAL_ENTIDAD", nombre_firmante="Representante legal de prueba",
            cargo_firmante="Representante legal", estado="FIRMADA",
        ),
    ])
    convenio.estado = "VIGENTE"
    convenio.fecha_firma = convenio.fecha_inicio
    convenio.etapa_actual = db.scalar(select(Etapa).where(Etapa.codigo == "FIRMA_ARCHIVO_SEGUIMIENTO"))
    db.commit()
    db.expire_all()
    return convenio


def _texto_pdf(contenido):
    """Comprueba estructura y lee los streams ASCII85/Flate de ReportLab.

    No cambia la compresión del generador ni introduce un parser en producción.
    Los literales Tj usan WinAnsi, con escapes PDF y caracteres octales.
    """
    assert contenido.startswith(b"%PDF-")
    fin = re.search(rb"startxref\s+(\d+)\s+%%EOF\s*$", contenido)
    assert fin is not None
    assert contenido[int(fin[1]):].startswith(b"xref")
    assert b"/Type /Catalog" in contenido
    assert b"/Type /Page" in contenido
    streams = re.findall(rb"stream\r?\n(.*?)endstream", contenido, re.DOTALL)
    assert streams
    textos = []
    for stream in streams:
        pagina = zlib.decompress(a85decode(b"<~" + stream.strip(), adobe=True))
        for literal in re.findall(rb"\(((?:\\.|[^\\)])*)\)\s*Tj", pagina):
            def desescapar(match):
                valor = match[1]
                if valor[:1] in b"01234567":
                    return bytes([int(valor, 8)])
                return {b"n": b"\n", b"r": b"\r", b"t": b"\t"}.get(valor, valor)

            textos.append(re.sub(rb"\\([0-7]{1,3}|.)", desescapar, literal).decode("cp1252"))
    assert textos
    return " ".join(" ".join(textos).split())


def _conteos(db):
    return {
        modelo.__tablename__: db.scalar(select(func.count()).select_from(modelo))
        for modelo in (Convenio, SolicitudConvenio, VersionConvenio, HistorialEtapa, Auditoria)
    }


@pytest.mark.parametrize("origen", ["MANUAL_ORI", "SOLICITUD"])
@pytest.mark.parametrize("modo", ["BORRADOR", "OFICIAL"])
@pytest.mark.parametrize("formato", ["pdf", "excel"])
def test_exporta_archivo_valido_con_datos_reales_y_sin_escrituras(
    client, db, request, origen, modo, formato,
):
    if origen == "MANUAL_ORI":
        fixture = "manual_documentado" if modo == "BORRADOR" else "manual_formalizado"
        convenio = request.getfixturevalue(fixture)
        assert convenio.solicitud_id is None and convenio.solicitud is None
        if modo == "OFICIAL":
            assert convenio.estado == "VIGENTE"
            assert convenio.etapa_actual.codigo == "FIRMA_ARCHIVO_SEGUIMIENTO"
        clausula = "Cláusula documental de auditoría"
    else:
        convenio = request.getfixturevalue("convenio_exportable")
        assert convenio.solicitud_id is not None
        convenio.solicitud.justificacion = "Justificación persistida de la solicitud"
        convenio.solicitud.contacto_contraparte_nombre = "Contacto de solicitud"
        convenio.solicitud.contacto_contraparte_correo = "contacto-solicitud@example.com"
        version = db.scalar(select(VersionConvenio).where(
            VersionConvenio.convenio_id == convenio.id,
            VersionConvenio.numero == convenio.version_actual,
        ))
        clausula = "Cláusula persistida del convenio desde solicitud"
        version.contenido = {"type": "doc", "content": [{"type": "paragraph", "content": [
            {"type": "text", "text": clausula},
        ]}]}
        if modo == "BORRADOR":
            convenio.etapa_actual = db.scalar(select(Etapa).where(Etapa.codigo == "ELABORACION"))
        db.commit()
        db.expire_all()
    antes = _conteos(db)
    dto = ServicioPdfConvenio().mapear_convenio_a_dto(convenio.id, db)
    assert dto.modo_emision.value == modo
    assert dto.objeto == convenio.objeto
    assert dto.aliado_nombre == convenio.aliado.nombre
    assert dto.aliado_identificacion == convenio.aliado.identificacion
    assert [seccion.texto for seccion in dto.secciones_contenido] == [clausula]
    assert (len(dto.firmantes) > 0) is (modo == "OFICIAL")
    if origen == "MANUAL_ORI":
        for campo in ("solicitud_consecutivo", "justificacion", "solicitante_tipo", "solicitante_nombre", "solicitante_correo"):
            assert getattr(dto, campo) == NO_APLICA
        assert dto.aliado_contacto_nombre == NO_DISPONIBLE
        assert dto.aliado_contacto_correo == convenio.aliado.correo
        assert all(firma.estado == "FIRMADA" for firma in dto.firmantes)
    else:
        assert dto.solicitud_consecutivo == convenio.solicitud.consecutivo
        assert dto.justificacion == convenio.solicitud.justificacion
        assert dto.solicitante_nombre == convenio.solicitud.solicitante.nombre_completo
        assert dto.solicitante_correo == convenio.solicitud.solicitante.correo
        assert dto.solicitante_tipo == convenio.solicitud.tipo_solicitante
        assert dto.aliado_contacto_nombre == "Contacto de solicitud"
        assert dto.aliado_contacto_correo == "contacto-solicitud@example.com"
    respuesta = client.get(f"/api/convenios/{convenio.id}/exportar", params={"formato": formato})
    assert respuesta.status_code == 200, respuesta.text
    prefijo = "BORRADOR_Convenio_" if modo == "BORRADOR" else "Convenio_"
    assert f'attachment; filename="{prefijo}' in respuesta.headers["content-disposition"]
    datos_reales = [convenio.objeto, convenio.aliado.nombre, convenio.aliado.identificacion, clausula]
    if formato == "pdf":
        assert respuesta.headers["content-type"] == "application/pdf"
        texto = _texto_pdf(respuesta.content)
        assert all(valor in texto for valor in datos_reales)
        assert ("BORRADOR - NO FORMALIZADO" in texto or "BORRADOR — NO FORMALIZADO" in texto) is (modo == "BORRADOR")
        assert ("DOCUMENTO OFICIAL DE CONVENIO" in texto) is (modo == "OFICIAL")
        if origen == "MANUAL_ORI":
            assert NO_APLICA in texto
        else:
            # La celda puede repartir un consecutivo largo entre varias líneas.
            assert convenio.solicitud.consecutivo in texto.replace(" ", "")
            assert convenio.solicitud.justificacion in texto
        for firma in dto.firmantes:
            assert firma.nombre in texto
    else:
        assert respuesta.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        libro = load_workbook(BytesIO(respuesta.content))
        assert libro.sheetnames == ["Datos Generales", "Participantes y Compromisos"]
        valores = {str(celda.value) for hoja in libro for fila in hoja for celda in fila if celda.value is not None}
        assert set(datos_reales) <= valores
        assert dto.solicitud_consecutivo in valores
        assert dto.justificacion in valores
        assert dto.aliado_contacto_correo in valores
        generales = {fila[0]: fila[1] for fila in libro["Datos Generales"].values if fila[0]}
        assert generales["Consecutivo de la solicitud"] == dto.solicitud_consecutivo
        assert generales["Fecha de inicio de vigencia"].date() == convenio.fecha_inicio
        assert generales["Fecha de fin de vigencia"].date() == convenio.fecha_vencimiento
        assert generales["Versión del documento"] == convenio.version_actual
        assert ("BORRADOR — NO FORMALIZADO" in valores) is (modo == "BORRADOR")
        assert {firma.nombre for firma in dto.firmantes} <= valores
    assert _conteos(db) == antes


@pytest.mark.parametrize("etapa", ["ELABORACION", "REVISION_JURIDICA"])
@pytest.mark.parametrize("formato", ["pdf", "excel"])
def test_manual_sin_aliado_reporta_validacion_controlada(
    client, db, manual_documentado, etapa, formato,
):
    convenio = manual_documentado
    convenio.aliado_id = None
    convenio.etapa_actual = db.scalar(select(Etapa).where(Etapa.codigo == etapa))
    db.commit()
    db.expire_all()
    antes = _conteos(db)
    respuesta = client.get(f"/api/convenios/{convenio.id}/exportar", params={"formato": formato})
    assert respuesta.status_code == 422, respuesta.text
    assert respuesta.json()["error"] == "CONVENIO_INCOMPLETO"
    esperados = ["entidad_participante"]
    if etapa == "REVISION_JURIDICA":
        esperados.extend(["representante_legal_institucion", "representante_legal_entidad"])
    assert respuesta.json()["campos_faltantes"] == esperados
    assert _conteos(db) == antes


@pytest.mark.parametrize("campo,valor,faltantes", [
    ("fecha_inicio", None, ["fecha_inicio_vigencia"]),
    ("fecha_vencimiento", None, ["fecha_fin_vigencia"]),
    ("tipo_convenio_id", None, ["tipo_convenio"]),
    ("objeto", "  ", ["objeto"]),
    ("version_actual", 999, ["version_contractual", "representante_legal_institucion", "representante_legal_entidad"]),
])
@pytest.mark.parametrize("formato", ["pdf", "excel"])
def test_manual_oficial_no_relaja_campos_obligatorios(
    client, db, manual_formalizado, campo, valor, faltantes, formato,
):
    convenio = manual_formalizado
    setattr(convenio, campo, valor)
    db.commit()
    db.expire_all()
    antes = _conteos(db)
    respuesta = client.get(f"/api/convenios/{convenio.id}/exportar", params={"formato": formato})
    assert respuesta.status_code == 422, respuesta.text
    assert respuesta.json()["error"] == "CONVENIO_INCOMPLETO"
    assert respuesta.json()["campos_faltantes"] == faltantes
    assert _conteos(db) == antes


@pytest.mark.parametrize("formato", ["pdf", "excel"])
@pytest.mark.parametrize("contenido,faltante", [
    ({"type": "doc", "content": [{"type": "paragraph"}]}, "clausulas"),
    ({"type": "html"}, "clausulas"),
])
def test_manual_no_exporta_documento_vacio_o_invalido(
    client, db, manual_documentado, formato, contenido, faltante,
):
    convenio = manual_documentado
    version = db.scalar(select(VersionConvenio).where(
        VersionConvenio.convenio_id == convenio.id, VersionConvenio.numero == convenio.version_actual,
    ))
    version.contenido = contenido
    db.commit()
    db.expire_all()
    respuesta = client.get(f"/api/convenios/{convenio.id}/exportar", params={"formato": formato})
    assert respuesta.status_code == 422, respuesta.text
    assert respuesta.json()["campos_faltantes"] == [faltante]


@pytest.mark.parametrize("etapa", ["ELABORACION", "REVISION_JURIDICA"])
def test_solicitud_ausente_es_error_de_integridad_solo_en_origen_solicitud(etapa):
    # La BD impide este registro inconsistente. El validador también lo indica
    # cuando recibe un objeto incompleto, en vez de considerarlo creación manual.
    convenio = Convenio(origen="SOLICITUD", estado="EN_TRAMITE", etapa_actual=Etapa(codigo=etapa))
    faltantes = ConvenioExportValidator().campos_faltantes(convenio)
    assert "solicitud_origen" in faltantes
    assert "entidad_participante" in faltantes


@pytest.mark.parametrize("prefijo", ["=", "+", "-", "@"])
def test_excel_manual_conserva_textos_peligrosos_como_literales(
    client, db, manual_documentado, prefijo,
):
    texto = f"{prefijo}SUM(A1:A2)"
    manual_documentado.objeto = texto
    manual_documentado.aliado.nombre = texto
    db.commit()
    respuesta = client.get(f"/api/convenios/{manual_documentado.id}/exportar?formato=excel")
    assert respuesta.status_code == 200, respuesta.text
    libro = load_workbook(BytesIO(respuesta.content))
    celdas = [celda for hoja in libro for fila in hoja for celda in fila if celda.value == f"'{texto}"]
    assert len(celdas) == 2
    assert all(celda.data_type == "s" for celda in celdas)

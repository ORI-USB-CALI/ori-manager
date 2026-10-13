from collections.abc import Callable
from datetime import date
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from backend.core.roles import CodigoRol, TipoUsuario
from backend.models.actividad_convenio import ActividadConvenio
from backend.models.convenio import Convenio
from backend.models.enums import (
    EstadoActividadConvenio,
    EstadoConvenio,
    EstadoSolicitud,
    TipoSolicitante,
)
from backend.models.rol import Rol
from backend.models.solicitud_convenio import SolicitudConvenio
from backend.models.usuario import Usuario
from backend.schemas.actividad_utilizacion import ActividadUtilizacionCrear
from backend.services import actividades_utilizacion
from backend.services.actividades_utilizacion import (
    ServicioActividadesUtilizacion,
)
from backend.services.convenios import ConvenioNoEditable
from backend.services.reportes import ServicioReportes

HOY = date(2026, 10, 13)
INICIO = date(2026, 1, 1)
VENCIMIENTO = date(2027, 1, 1)
URL_REPORTE = "/api/convenios/reportes/activos"

ACTIVIDAD_EJEMPLO = {
    "fecha_ejecucion": "2026-09-15",
    "titulo": "Movilidad academica de dos estudiantes de Ingenieria",
    "descripcion": "Participación en intercambio académico durante el semestre 2026-2",
    "responsable_nombre": "Facultad de Ingeniería",
    "observaciones": "Estudiantes seleccionados por convocatoria interna",
}


def url_actividades(convenio_id: int) -> str:
    return f"/api/convenios/{convenio_id}/actividades-utilizacion"


@pytest.fixture(autouse=True)
def fecha_dominio_fija(monkeypatch) -> None:
    monkeypatch.setattr(actividades_utilizacion, "fecha_actual_dominio", lambda: HOY)


@pytest.fixture
def crear_convenio_en_estado(
    db: Session, gestor: Usuario, crear_convenio
) -> Callable[..., Convenio]:
    def _crear(
        estado: EstadoConvenio = EstadoConvenio.VIGENTE,
        fecha_inicio: date | None = INICIO,
        fecha_vencimiento: date | None = VENCIMIENTO,
    ) -> Convenio:
        convenio = crear_convenio(
            gestor,
            codigo=f"HU35-{uuid4().hex[:12]}",
            fecha_inicio=fecha_inicio,
            fecha_vencimiento=fecha_vencimiento,
        )
        convenio.estado = estado.value
        db.commit()
        db.refresh(convenio)
        return convenio

    return _crear


def actividades_guardadas(db: Session) -> list[ActividadConvenio]:
    return list(db.scalars(select(ActividadConvenio)))


# ---- CA-01: registrar una actividad -----------------------------------------


@pytest.mark.parametrize(
    "estado", [EstadoConvenio.VIGENTE, EstadoConvenio.POR_VENCER]
)
def test_ca01_registra_actividad_asociada_al_convenio(
    client, db: Session, gestor, crear_convenio_en_estado, estado
) -> None:
    convenio = crear_convenio_en_estado(estado)

    respuesta = client.post(url_actividades(convenio.id), json=ACTIVIDAD_EJEMPLO)

    assert respuesta.status_code == 201
    assert respuesta.json()["convenio_id"] == convenio.id
    guardada = db.get(ActividadConvenio, respuesta.json()["id"])
    assert guardada is not None
    assert guardada.convenio_id == convenio.id
    assert guardada.estado == EstadoActividadConvenio.REALIZADA.value


# ---- CA-02: conservar la información registrada -----------------------------


def test_ca02_conserva_los_datos_registrados(
    client, db: Session, gestor, crear_convenio_en_estado
) -> None:
    convenio = crear_convenio_en_estado()

    cuerpo = client.post(
        url_actividades(convenio.id), json=ACTIVIDAD_EJEMPLO
    ).json()

    guardada = db.get(ActividadConvenio, cuerpo["id"])
    assert guardada.fecha_ejecucion == date(2026, 9, 15)
    assert guardada.titulo == ACTIVIDAD_EJEMPLO["titulo"]
    assert guardada.descripcion == ACTIVIDAD_EJEMPLO["descripcion"]
    assert guardada.responsable_nombre == ACTIVIDAD_EJEMPLO["responsable_nombre"]
    assert guardada.observaciones == ACTIVIDAD_EJEMPLO["observaciones"]
    assert guardada.responsable_id is None
    assert guardada.registrado_por_id == gestor.id
    assert guardada.creado_en is not None
    assert cuerpo["registrado_por_id"] == gestor.id
    assert cuerpo["estado"] == EstadoActividadConvenio.REALIZADA.value


@pytest.mark.parametrize("observaciones", [None, "", "   "])
def test_ca02_observaciones_es_opcional(
    client, gestor, crear_convenio_en_estado, observaciones
) -> None:
    convenio = crear_convenio_en_estado()
    datos = {**ACTIVIDAD_EJEMPLO, "observaciones": observaciones}

    respuesta = client.post(url_actividades(convenio.id), json=datos)

    assert respuesta.status_code == 201
    assert respuesta.json()["observaciones"] is None


@pytest.mark.parametrize(
    "campo", ["fecha_ejecucion", "titulo", "descripcion", "responsable_nombre"]
)
def test_ca02_rechaza_campos_obligatorios_faltantes(
    client, db: Session, gestor, crear_convenio_en_estado, campo
) -> None:
    convenio = crear_convenio_en_estado()
    datos = {k: v for k, v in ACTIVIDAD_EJEMPLO.items() if k != campo}

    respuesta = client.post(url_actividades(convenio.id), json=datos)

    assert respuesta.status_code == 422
    assert actividades_guardadas(db) == []


@pytest.mark.parametrize("campo", ["titulo", "descripcion", "responsable_nombre"])
def test_ca02_rechaza_textos_obligatorios_vacios(
    client, gestor, crear_convenio_en_estado, campo
) -> None:
    convenio = crear_convenio_en_estado()
    datos = {**ACTIVIDAD_EJEMPLO, campo: "   "}

    respuesta = client.post(url_actividades(convenio.id), json=datos)

    assert respuesta.status_code == 422


# ---- CA-02: fecha dentro de la vigencia y no futura -------------------------


@pytest.mark.parametrize(
    "fecha", [INICIO.isoformat(), HOY.isoformat()], ids=["inicio", "hoy"]
)
def test_fecha_acepta_limites_inicio_y_hoy(
    client, gestor, crear_convenio_en_estado, fecha
) -> None:
    convenio = crear_convenio_en_estado()

    respuesta = client.post(
        url_actividades(convenio.id), json={**ACTIVIDAD_EJEMPLO, "fecha_ejecucion": fecha}
    )

    assert respuesta.status_code == 201


def test_fecha_acepta_limite_de_vencimiento(
    client, gestor, crear_convenio_en_estado
) -> None:
    convenio = crear_convenio_en_estado(fecha_vencimiento=date(2026, 9, 30))

    respuesta = client.post(
        url_actividades(convenio.id),
        json={**ACTIVIDAD_EJEMPLO, "fecha_ejecucion": "2026-09-30"},
    )

    assert respuesta.status_code == 201


@pytest.mark.parametrize(
    ("fecha_vencimiento", "fecha", "mensaje"),
    [
        (VENCIMIENTO, "2025-12-31", "período de vigencia"),
        (date(2026, 9, 30), "2026-10-01", "período de vigencia"),
        (VENCIMIENTO, "2026-10-14", "futura"),
    ],
    ids=["antes-del-inicio", "despues-del-vencimiento", "futura"],
)
def test_fecha_rechaza_fuera_del_periodo_o_futura(
    client,
    db: Session,
    gestor,
    crear_convenio_en_estado,
    fecha_vencimiento,
    fecha,
    mensaje,
) -> None:
    convenio = crear_convenio_en_estado(fecha_vencimiento=fecha_vencimiento)

    respuesta = client.post(
        url_actividades(convenio.id), json={**ACTIVIDAD_EJEMPLO, "fecha_ejecucion": fecha}
    )

    assert respuesta.status_code == 422
    assert mensaje in respuesta.json()["detail"]
    assert actividades_guardadas(db) == []


@pytest.mark.parametrize(
    ("fecha_inicio", "fecha_vencimiento"),
    [(None, VENCIMIENTO), (INICIO, None), (None, None)],
    ids=["sin-inicio", "sin-vencimiento", "sin-ambas"],
)
def test_rechaza_convenio_sin_periodo_de_vigencia(
    client,
    db: Session,
    gestor,
    crear_convenio_en_estado,
    fecha_inicio,
    fecha_vencimiento,
) -> None:
    convenio = crear_convenio_en_estado(
        fecha_inicio=fecha_inicio, fecha_vencimiento=fecha_vencimiento
    )

    respuesta = client.post(url_actividades(convenio.id), json=ACTIVIDAD_EJEMPLO)

    assert respuesta.status_code == 409
    assert "período de vigencia" in respuesta.json()["detail"]
    assert actividades_guardadas(db) == []


# ---- CA-03: consultar actividades del convenio ------------------------------


def test_ca03_lista_actividades_del_convenio_ordenadas(
    client, gestor, crear_convenio_en_estado
) -> None:
    convenio = crear_convenio_en_estado()
    otro = crear_convenio_en_estado()
    antigua = client.post(
        url_actividades(convenio.id),
        json={**ACTIVIDAD_EJEMPLO, "fecha_ejecucion": "2026-03-01"},
    ).json()
    reciente = client.post(url_actividades(convenio.id), json=ACTIVIDAD_EJEMPLO).json()
    mismo_dia = client.post(url_actividades(convenio.id), json=ACTIVIDAD_EJEMPLO).json()
    client.post(url_actividades(otro.id), json=ACTIVIDAD_EJEMPLO)

    respuesta = client.get(url_actividades(convenio.id))

    assert respuesta.status_code == 200
    assert [item["id"] for item in respuesta.json()] == [
        mismo_dia["id"],
        reciente["id"],
        antigua["id"],
    ]
    assert respuesta.json()[1] == reciente


def test_ca03_no_mezcla_actividades_planificadas(
    client, db: Session, gestor, crear_convenio_en_estado
) -> None:
    convenio = crear_convenio_en_estado()
    realizada = client.post(url_actividades(convenio.id), json=ACTIVIDAD_EJEMPLO).json()
    db.add(
        ActividadConvenio(
            convenio_id=convenio.id,
            titulo="Próximo seminario conjunto",
            descripcion="Actividad planificada, no ejecutada",
            fecha_planificada=date(2026, 11, 20),
            estado=EstadoActividadConvenio.PENDIENTE.value,
            registrado_por_id=gestor.id,
        )
    )
    db.commit()

    respuesta = client.get(url_actividades(convenio.id))

    assert [item["id"] for item in respuesta.json()] == [realizada["id"]]


# ---- CA-04: restringir el registro ------------------------------------------


@pytest.mark.parametrize(
    ("codigo_rol", "tipo_usuario"),
    [
        (CodigoRol.REVISOR_ORI, TipoUsuario.INTERNO),
        (CodigoRol.SOLICITANTE_INTERNO, TipoUsuario.INTERNO),
        (CodigoRol.SOLICITANTE_EXTERNO, TipoUsuario.EXTERNO),
    ],
)
def test_ca04_rol_sin_permiso_no_puede_registrar(
    client,
    db: Session,
    crear_usuario,
    entrar_como,
    crear_convenio_en_estado,
    codigo_rol,
    tipo_usuario,
) -> None:
    convenio = crear_convenio_en_estado()
    entrar_como(crear_usuario(codigo_rol, tipo_usuario))

    respuesta = client.post(url_actividades(convenio.id), json=ACTIVIDAD_EJEMPLO)

    assert respuesta.status_code == 403
    assert actividades_guardadas(db) == []


def test_ca04_sin_sesion_no_puede_registrar(
    client, gestor, crear_convenio_en_estado
) -> None:
    convenio = crear_convenio_en_estado()
    client.cookies.clear()

    respuesta = client.post(url_actividades(convenio.id), json=ACTIVIDAD_EJEMPLO)

    assert respuesta.status_code == 401


def test_ca04_administrador_puede_registrar(
    client, crear_usuario, entrar_como, crear_convenio_en_estado
) -> None:
    convenio = crear_convenio_en_estado()
    administrador = crear_usuario(CodigoRol.ADMINISTRADOR_ORI, TipoUsuario.INTERNO)
    entrar_como(administrador)

    respuesta = client.post(url_actividades(convenio.id), json=ACTIVIDAD_EJEMPLO)

    assert respuesta.status_code == 201
    assert respuesta.json()["registrado_por_id"] == administrador.id


@pytest.mark.parametrize(
    "estado",
    [
        EstadoConvenio.EN_TRAMITE,
        EstadoConvenio.VENCIDO,
        EstadoConvenio.RENOVADO,
        EstadoConvenio.FINALIZADO,
        EstadoConvenio.CANCELADO,
    ],
)
def test_ca04_rechaza_registro_en_convenio_no_activo(
    client, db: Session, gestor, crear_convenio_en_estado, estado
) -> None:
    convenio = crear_convenio_en_estado(estado)

    respuesta = client.post(url_actividades(convenio.id), json=ACTIVIDAD_EJEMPLO)

    assert respuesta.status_code == 409
    assert "vigentes o por vencer" in respuesta.json()["detail"]
    assert actividades_guardadas(db) == []


@pytest.mark.parametrize(
    "estado", [EstadoConvenio.VIGENTE, EstadoConvenio.POR_VENCER]
)
def test_registro_recarga_estado_confirmado_por_otra_conexion(
    db_engine, estado
) -> None:
    # No usar db: sus commits solo liberan savepoints de una transacción externa.
    assert db_engine.dialect.name == "postgresql"
    with Session(db_engine) as preparacion, preparacion.begin():
        rol = preparacion.scalar(
            select(Rol).where(Rol.codigo == CodigoRol.GESTOR_ORI.value)
        )
        assert rol is not None
        usuario = Usuario(
            correo=f"hu35-concurrencia-{uuid4().hex}@example.com",
            hash_contrasena="sin-login",
            nombre_completo="Gestor de prueba HU-35",
            rol=rol,
            tipo_usuario=TipoUsuario.INTERNO.value,
        )
        preparacion.add(usuario)
        preparacion.flush()
        solicitud = SolicitudConvenio(
            consecutivo=f"SOL-{uuid4().hex}",
            solicitante_id=usuario.id,
            tipo_solicitante=TipoSolicitante.INTERNO.value,
            estado=EstadoSolicitud.APROBADA.value,
        )
        preparacion.add(solicitud)
        preparacion.flush()
        convenio = Convenio(
            solicitud_id=solicitud.id,
            creado_por_id=usuario.id,
            estado=estado.value,
            fecha_inicio=INICIO,
            fecha_vencimiento=VENCIMIENTO,
        )
        preparacion.add(convenio)
        preparacion.flush()
        convenio_id, solicitud_id, usuario_id = convenio.id, solicitud.id, usuario.id
    # La salida de begin() confirma los datos para ambas conexiones.
    conteo = (
        select(func.count())
        .select_from(ActividadConvenio)
        .where(ActividadConvenio.convenio_id == convenio_id)
    )
    try:
        with Session(db_engine) as registro, Session(db_engine) as cambio:
            # Ambas sesiones mantienen conexiones abiertas al comparar los PID.
            assert registro.scalar(select(func.pg_backend_pid())) != cambio.scalar(
                select(func.pg_backend_pid())
            )
            precargado = registro.get(Convenio, convenio_id)
            assert precargado is not None
            assert precargado.estado == estado.value
            gestor = registro.get(Usuario, usuario_id)
            assert gestor is not None
            assert registro.scalar(conteo) == 0

            actualizado = cambio.get(Convenio, convenio_id)
            assert actualizado is not None
            assert actualizado.estado == estado.value
            actualizado.estado = EstadoConvenio.FINALIZADO.value
            cambio.commit()

            # Conservar una referencia fuerte evita que el identity map la pierda.
            assert registro.get(Convenio, convenio_id) is precargado
            assert precargado.estado == estado.value
            with pytest.raises(ConvenioNoEditable, match="vigentes o por vencer"):
                ServicioActividadesUtilizacion(registro).registrar(
                    convenio_id,
                    ActividadUtilizacionCrear.model_validate(ACTIVIDAD_EJEMPLO),
                    gestor,
                )
            assert precargado.estado == EstadoConvenio.FINALIZADO.value
            assert registro.scalar(conteo) == 0
        with Session(db_engine) as verificacion:
            assert verificacion.get(Convenio, convenio_id).estado == (
                EstadoConvenio.FINALIZADO.value
            )
            assert verificacion.scalar(conteo) == 0
    finally:
        # Limpiar también si la regresión inserta y confirma una actividad.
        with Session(db_engine) as limpieza, limpieza.begin():
            limpieza.execute(
                delete(ActividadConvenio).where(
                    ActividadConvenio.convenio_id == convenio_id
                )
            )
            limpieza.execute(delete(Convenio).where(Convenio.id == convenio_id))
            limpieza.execute(
                delete(SolicitudConvenio).where(SolicitudConvenio.id == solicitud_id)
            )
            limpieza.execute(delete(Usuario).where(Usuario.id == usuario_id))


def test_convenio_inexistente_responde_404(client, gestor) -> None:
    assert client.post(url_actividades(999_999), json=ACTIVIDAD_EJEMPLO).status_code == 404
    assert client.get(url_actividades(999_999)).status_code == 404


# ---- CA-05: disponibles para consumo externo --------------------------------


def test_ca05_convenio_sin_actividades_devuelve_lista_vacia(
    client, gestor, crear_convenio_en_estado
) -> None:
    convenio = crear_convenio_en_estado()

    respuesta = client.get(url_actividades(convenio.id))

    assert respuesta.status_code == 200
    assert respuesta.json() == []


def test_ca05_listado_refleja_registro_inmediatamente(
    client, gestor, crear_convenio_en_estado
) -> None:
    convenio = crear_convenio_en_estado()
    assert client.get(url_actividades(convenio.id)).json() == []

    creada = client.post(url_actividades(convenio.id), json=ACTIVIDAD_EJEMPLO).json()

    assert client.get(url_actividades(convenio.id)).json() == [creada]


def test_ca05_servicio_lista_actividades_sin_volver_a_registrarlas(
    client, db: Session, gestor, crear_convenio_en_estado
) -> None:
    convenio = crear_convenio_en_estado()
    creada = client.post(url_actividades(convenio.id), json=ACTIVIDAD_EJEMPLO).json()

    actividades = ServicioActividadesUtilizacion(db).listar(convenio.id, gestor)

    assert [actividad.id for actividad in actividades] == [creada["id"]]


def test_ca05_renovacion_no_hereda_actividades_del_origen(
    client, db: Session, gestor, crear_convenio_en_estado
) -> None:
    origen = crear_convenio_en_estado()
    client.post(url_actividades(origen.id), json=ACTIVIDAD_EJEMPLO)
    renovacion = crear_convenio_en_estado()
    renovacion.convenio_origen_id = origen.id
    db.commit()

    assert len(client.get(url_actividades(origen.id)).json()) == 1
    assert client.get(url_actividades(renovacion.id)).json() == []


# ---- Integración HU-35 → HU-34 ----------------------------------------------


def test_integracion_reporte_consulta_la_misma_actividad_registrada(
    client, db: Session, gestor, crear_convenio_en_estado
) -> None:
    convenio = crear_convenio_en_estado()
    creada = client.post(url_actividades(convenio.id), json=ACTIVIDAD_EJEMPLO).json()

    reporte = ServicioReportes(db).listar_convenios_activos()

    del_convenio = next(item for item in reporte if item.id == convenio.id)
    assert [actividad.id for actividad in del_convenio.actividades] == [creada["id"]]
    respuesta = client.get(URL_REPORTE)
    assert respuesta.status_code == 200
    en_api = next(item for item in respuesta.json() if item["id"] == convenio.id)
    assert [actividad["id"] for actividad in en_api["actividades"]] == [creada["id"]]
    assert en_api["actividades"][0]["fecha_ejecucion"] == ACTIVIDAD_EJEMPLO["fecha_ejecucion"]


def test_integracion_actividad_aparece_en_pdf_y_excel(
    client, gestor, crear_convenio_en_estado
) -> None:
    convenio = crear_convenio_en_estado()
    client.post(url_actividades(convenio.id), json=ACTIVIDAD_EJEMPLO)

    pdf = client.get(f"{URL_REPORTE}/pdf")
    excel = client.get(f"{URL_REPORTE}/excel")

    assert pdf.status_code == 200
    # El título se parte en varias líneas dentro de la celda del PDF.
    assert b"Movilidad" in pdf.content
    assert excel.status_code == 200
    hoja = load_workbook(BytesIO(excel.content))["Actividades"]
    titulos = [fila[1] for fila in hoja.iter_rows(min_row=2, values_only=True)]
    assert ACTIVIDAD_EJEMPLO["titulo"] in titulos


def test_integracion_convenio_sin_actividades_sigue_en_el_reporte(
    client, gestor, crear_convenio_en_estado
) -> None:
    convenio = crear_convenio_en_estado()

    respuesta = client.get(URL_REPORTE)

    en_api = next(item for item in respuesta.json() if item["id"] == convenio.id)
    assert en_api["actividades"] == []


def test_integracion_actividad_sigue_disponible_cuando_el_convenio_finaliza(
    client, db: Session, gestor, crear_convenio_en_estado
) -> None:
    convenio = crear_convenio_en_estado()
    creada = client.post(url_actividades(convenio.id), json=ACTIVIDAD_EJEMPLO).json()
    convenio.estado = EstadoConvenio.FINALIZADO.value
    db.commit()

    respuesta = client.get(url_actividades(convenio.id))

    assert respuesta.status_code == 200
    assert respuesta.json() == [creada]
    assert db.get(ActividadConvenio, creada["id"]) is not None
